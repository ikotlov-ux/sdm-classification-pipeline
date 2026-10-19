#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
09_MatrixLegend.py — матричная легенда для полигонального слоя.

Класс полигона описывается двумя векторами (например, тип растительности ×
стадия сукцессии). Скрипт строит матрицу «вектор по вертикали × вектор по
горизонтали», раскрашивает её и сохраняет в XLSX (цвета — HEX).

Порядок работы в окне
---------------------
1. Слой полигонов: файл и слой.
2. Где записаны векторы:
   • в двух полях — выбрать поле по горизонтали и поле по вертикали;
   • в одном поле — выбрать поле и способ разбора кода:
       - «первые k цифр | остальные»: 3125 при k = 1 → 3 и 125;
       - «остальные | последние k цифр»: 12125 при k = 3 → 12 и 125;
       - «разделитель»: 3_125, 3-125, 3.125 → 3 и 125;
     и указать, по какой оси идёт первая часть кода.
3. Названия классов: CSV с двумя столбцами «номер; название» для каждого
   вектора (разделитель и кодировка UTF-8 / Windows-1251 определяются сами,
   строка заголовка допускается). Можно один файл на оба вектора.
4. «Построить матрицу». Затем:
   • какая ось задаётся цветом (по умолчанию — ось с бо́льшим числом классов),
     другая — тоном;
   • цвет каждой категории цветовой оси (автопалитра + выбор вручную);
   • способ изменения тона (по умолчанию для всех категорий):
       - к бледному (смешение с белым);
       - к тёмному (смешение с чёрным);
       - к серому (снижение насыщенности);
       - от бледного через базовый к тёмному;
       - светлота HLS (оттенок и насыщенность сохраняются);
       - случайная в пределах цвета (оттенок ± немного, разная светлота);
     с силой, обратным порядком и зерном для случайной шкалы;
   • для любой категории можно задать свой градиент (кнопка в списке);
   • любую ячейку можно перекрасить вручную.
5. Перекодировка типов (раздел 4 окна): «Старый код → Новый код», «Выполнить».
   Слияние: 1400 → 3400 — полигоны ячейки 1400 добавляются к 3400; перенумерация:
   1400 → 1450 — ячейка получает новый код (новая строка / столбец появятся сами,
   пустые исчезнут). Число полигонов и площадь пересчитываются сразу.
   Два поля: ячейка задаётся парой «строка столбец» (1 400, 1;400, 1|400).
   «Отменить» / «Повторить» (Ctrl+Z / Ctrl+Y) — до 50 шагов назад.
   ПКМ по ячейке → «старый код = …» / «новый код = …» подставляет её код.
   Сам слой не меняется: в стилях QGIS / ArcGIS исходные коды (1400, 2400…)
   получают цвет и подпись новой ячейки. По флажку «Записать перекодированный
   слой» сохраняется копия слоя с новыми кодами (<слой>_legend_recoded.gpkg).
6. «Сохранить XLSX».

Матрица на экране
-----------------
    ЛКМ по ячейке           — изменить цвет ячейки
    ЛКМ по заголовку цвета  — изменить базовый цвет категории
    ПКМ                     — меню (HEX, сброс, градиент категории)
    перетаскивание ЛКМ/СКМ  — сдвиг; стрелки — сдвиг
    колесо мыши             — масштаб к точке под курсором
    Shift + колесо          — сдвиг по горизонтали
    Home                    — вписать матрицу в окно
Заголовки строк и столбцов закреплены.

Выход: 09_MATRIX_LEGEND/<ддммгг_ччмм><_vXXX источника>/<слой>_legend.xlsx
    лист «Матрица»   — раскрашенная матрица; в ячейке через горизонтальную черту:
                       число полигонов (сверху) и суммарная площадь (снизу);
                       подписи строк и столбцов — названия классов из CSV
    лист «Классы»    — все сочетания из слоя: код, значения, названия, цвет HEX, RGB,
                       число полигонов, площадь
    лист «Перекодировка» — шаги «старый → новый», перенесённые объекты и площадь
                       (при перекодировке; в «Классах» — столбец «Объединено из»)
    лист «Цвета»     — базовые цвета и градиенты категорий цветовой оси
    лист «Параметры» — исходные данные и настройки (по ним раскраску можно
                       загрузить обратно кнопкой «Загрузить раскраску»;
                       перекодировки восстанавливаются вместе с ней)
    <слой>_legend_qgis.qml       — стиль QGIS (Свойства слоя → Стиль → Загрузить стиль)
    <слой>_legend_arcgispro.lyrx — ArcGIS Pro (Символы → Импорт / «Применить символы из слоя»)
    <слой>_legend_arcmap108.py   — ArcMap 10.8: выполнить в окне Python ArcMap
                                   (execfile), слой раскрасится и сохранится
                                   <слой>_legend_arcmap108.lyr
    Подписи классов в стилях: «строка — столбец». Два поля: QGIS — выражение
    concat("строка", '|', "столбец"); ArcGIS — уникальные значения по двум полям.

Запуск из папки проекта:
    %run 09_MatrixLegend.py      # Jupyter — откроется окно
    python 09_MatrixLegend.py    # терминал — откроется окно

CLI (раскраска по умолчанию, без окна):
    python 09_MatrixLegend.py map.gpkg --layer classes --field code --parse first:1 \\
        --first-axis v --names-v veg.csv --names-h stage.csv --out legend.xlsx
    python 09_MatrixLegend.py map.gpkg --fields STAGE VEG --tone tint
    python 09_MatrixLegend.py map.gpkg --field class_value --parse last:3 \\
        --recode 1400=3400 2400=3400 4400=3400 --write-layer

Зависимости:
    pip install numpy pandas geopandas pyogrio openpyxl
"""

from __future__ import annotations

import argparse
import colorsys
import copy
import csv
import io
import json
import math
import os
import random
import re
import sys
from datetime import datetime

try:
    sys.stdout.reconfigure(line_buffering=True)
except Exception:
    pass

OUT_BASE_NAME = "09_MATRIX_LEGEND"
VECTOR_TYPES = [("Векторные слои", "*.gpkg *.shp *.geojson *.json *.fgb *.gml *.kml *.sqlite *.tab"),
                ("Все файлы", "*.*")]
CSV_TYPES = [("CSV", "*.csv *.txt"), ("Все файлы", "*.*")]

PARSE_METHODS = [
    ("first", "Первые k цифр | остальные (3125 → 3 | 125)"),
    ("last", "Остальные | последние k цифр (12125 → 12 | 125)"),
    ("sep", "Разделитель (3_125 → 3 | 125)"),
]

TONE_MODES = [
    ("tint", "К бледному (смешение с белым)"),
    ("shade", "К тёмному (смешение с чёрным)"),
    ("grey", "К серому (снижение насыщенности)"),
    ("tint_shade", "От бледного через базовый к тёмному"),
    ("lightness", "Светлота HLS (насыщенность сохраняется)"),
    ("random", "Случайная в пределах цвета"),
]
TONE_LABEL = dict(TONE_MODES)
TONE_SHORT = {"tint": "к бледному", "shade": "к тёмному", "grey": "к серому",
              "tint_shade": "бледн.→тёмн.", "lightness": "светлота", "random": "случайная"}

PALETTES = {
    "Контрастная качественная": [
        "#e6194b", "#3cb44b", "#4363d8", "#f58231", "#911eb4", "#42d4f4",
        "#f032e6", "#bfef45", "#469990", "#9a6324", "#800000", "#808000",
        "#000075", "#ffe119", "#a9a9a9"],
    "Радуга (равные шаги оттенка)": None,
    "Золотое сечение (оттенок)": None,
    "Природная (зелёные, бурые, синие)": [
        "#1b7837", "#5aae61", "#a6611a", "#dfc27d", "#2166ac", "#67a9cf",
        "#762a83", "#c51b7d", "#8c510a", "#35978f", "#b2182b", "#7f3b08"],
}
DEFAULT_PALETTE = "Контрастная качественная"

AREA_FIELD = "__area_m2"
AREA_UNITS = {"га": 1e4, "км²": 1e6, "м²": 1.0}
AREA_FMT = {"га": "#,##0.00", "км²": "#,##0.000", "м²": "#,##0"}
LABEL_MODES = [("name", "Название"), ("num_name", "Номер — название"), ("num", "Номер")]
CELL_TEXT = [("count_area", "Число / площадь"), ("none", "Без подписи")]


def log(msg=""):
    print(msg, flush=True)


# ===========================================================================
# 1. ДАННЫЕ: ключи, разбор кодов, названия
# ===========================================================================

def norm_key(v):
    """Значение класса → строковый ключ: 3, 3.0, '3', '003' → '3'."""
    if v is None:
        return None
    try:
        import numpy as np
        if isinstance(v, np.generic):
            v = v.item()
    except Exception:
        pass
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if math.isnan(v):
            return None
        return str(int(v)) if v.is_integer() else repr(v)
    s = str(v).strip()
    if s == "" or s.lower() in ("nan", "none", "null"):
        return None
    if re.fullmatch(r"[+-]?\d+(\.0*)?", s):
        return str(int(float(s)))
    return s


def orig_str(v):
    """Значение поля как строка, совпадающая с тем, как его покажут ГИС (3.0 → '3')."""
    try:
        import numpy as np
        if isinstance(v, np.generic):
            v = v.item()
    except Exception:
        pass
    if isinstance(v, bool):
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    return str(v)


def _field_kind(series):
    import pandas as pd
    if pd.api.types.is_bool_dtype(series) or pd.api.types.is_integer_dtype(series):
        return "integer"
    if pd.api.types.is_float_dtype(series):
        return "double"
    return "string"


def raw_code(v):
    """Код как строка для разбора (3125.0 → '3125')."""
    k = norm_key(v)
    if k is None:
        return None
    # '003125' → разбирать как записано, если строка из цифр с ведущими нулями
    s = str(v).strip() if not isinstance(v, (int, float)) else k
    if re.fullmatch(r"\d+", s):
        return s
    return k


def parse_code(v, method, k=1, sep="_"):
    """Код из одного поля → (часть 1, часть 2) или None."""
    s = raw_code(v)
    if s is None:
        return None
    if method == "first":
        if len(s) <= k:
            return None
        a, b = s[:k], s[k:]
    elif method == "last":
        if len(s) <= k:
            return None
        a, b = s[:-k], s[-k:]
    elif method == "sep":
        if not sep or sep not in s:
            return None
        a, b = s.split(sep, 1)
    else:
        raise ValueError(f"Неизвестный способ разбора: {method}")
    a, b = norm_key(a), norm_key(b)
    if a is None or b is None:
        return None
    return a, b


UNDO_LIMIT = 50            # глубина отмены перекодировок


def sort_keys(keys):
    def k(x):
        try:
            return (0, float(x), "")
        except ValueError:
            return (1, 0.0, x.lower())
    return sorted(keys, key=k)


def read_names_csv(path):
    """CSV «номер; название» → {ключ: название}."""
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = [ln for ln in text.splitlines() if ln.strip()]
    if not lines:
        return {}
    sample = "\n".join(lines[:20])
    counts = {d: sample.count(d) for d in (";", "\t", ",", "|")}
    delim = max(counts, key=counts.get) if max(counts.values()) > 0 else ";"
    rows = list(csv.reader(io.StringIO("\n".join(lines)), delimiter=delim))
    out = {}
    for i, r in enumerate(rows):
        if len(r) < 2:
            continue
        key = norm_key(r[0])
        name = r[1].strip()
        if key is None:
            continue
        if i == 0 and not re.fullmatch(r"[+-]?\d+(\.\d+)?", r[0].strip()) and len(rows) > 1 \
                and re.fullmatch(r"[+-]?\d+(\.\d+)?", rows[1][0].strip() if rows[1] else ""):
            continue    # строка заголовка
        out[key] = name
    return out


def list_layers(path):
    try:
        import pyogrio
        return [str(x[0]) for x in pyogrio.list_layers(path)]
    except Exception:
        pass
    try:
        import fiona
        return list(fiona.listlayers(path))
    except Exception:
        return [os.path.splitext(os.path.basename(path))[0]]


def read_attrs(path, layer=None):
    """Атрибуты слоя + площадь каждого объекта в м² (колонка __area_m2).
    Географическая СК → площадь считается в UTM-зоне по центру данных."""
    import geopandas as gpd
    import pandas as pd
    kw = {}
    if layer and (len(list_layers(path)) > 1 or path.lower().endswith((".gpkg", ".sqlite", ".gdb"))):
        kw["layer"] = layer
    gdf = gpd.read_file(path, **kw)
    note = ""
    if "geometry" in gdf.columns and gdf.geometry.notna().any():
        g = gdf.geometry
        if gdf.crs is None:
            note = "СК слоя не задана — площадь в единицах координат слоя."
        elif gdf.crs.is_geographic:
            try:
                g = gdf.to_crs(gdf.estimate_utm_crs()).geometry
                note = "СК географическая — площадь рассчитана в UTM."
            except Exception:
                g = gdf.to_crs(6933).geometry
                note = "СК географическая — площадь рассчитана в EPSG:6933."
        area = g.area.fillna(0.0).values
    else:
        area = [0.0] * len(gdf)
        note = "В слое нет геометрии — площадь не рассчитана."
    df = pd.DataFrame(gdf.drop(columns=["geometry"], errors="ignore"))
    df[AREA_FIELD] = area
    df.attrs["area_note"] = note
    return df


# ===========================================================================
# 2. МОДЕЛЬ МАТРИЦЫ
# ===========================================================================

class MatrixModel:
    """Матрица: строки — вертикальная ось (v), столбцы — горизонтальная (h)."""

    def __init__(self):
        self.h_keys, self.v_keys = [], []
        self.h_names, self.v_names = {}, {}
        self.h_title, self.v_title = "по горизонтали", "по вертикали"
        self.counts = {}        # (hk, vk) -> число объектов
        self.areas = {}         # (hk, vk) -> площадь, м²
        self.label_mode = "name"
        self.area_unit = "га"
        self.orig = {}          # (hk, vk) -> исходные значения полей (для файлов стилей)
        self.field_types = {}   # поле -> integer | double | string
        self.codes = {}         # (hk, vk) -> исходные коды (одно поле)
        self.skipped = {}       # нераспознанные значения -> число
        self.source = {}
        self.hue_side = "h"
        self.base = {}          # категория цветовой оси -> HEX
        self.tone = {"mode": "tint", "strength": 0.75, "reverse": False, "seed": 1}
        self.override = {}      # категория цветовой оси -> {mode, strength, reverse}
        self.manual = {}        # (hk, vk) -> HEX
        self.hide_absent = True
        self.palette = DEFAULT_PALETTE
        # перекодировка типов: исходные сведения слоя + список шагов
        self._base = None
        self.recodes = []       # [{old, new, old_txt, new_txt, new_vals, n, area}]
        self.redo_stack = []
        self.undo_floor = 0     # шаги до этого индекса отменить нельзя (глубина UNDO_LIMIT)
        self.final = {}         # исходный кортеж значений -> новый кортеж (для перекодированного слоя)
        self.cell_code = {}     # ячейка -> код, введённый как «новый» (одно поле)

    # ---------- построение ----------
    @classmethod
    def from_table(cls, df, cfg):
        """cfg: mode two|one; h_field, v_field | field, method, k, sep, first_axis."""
        m = cls()
        m.source = dict(cfg)
        if AREA_FIELD not in df.columns:
            df = df.assign(**{AREA_FIELD: 0.0})
        m.field_types = {c: _field_kind(df[c]) for c in df.columns if c != AREA_FIELD}
        if cfg["mode"] == "two":
            hf, vf = cfg["h_field"], cfg["v_field"]
            m.h_title, m.v_title = hf, vf
            agg = df.groupby([hf, vf], dropna=False)[AREA_FIELD].agg(["size", "sum"])
            for (hv, vv), (n, a) in agg.iterrows():
                hk, vk = norm_key(hv), norm_key(vv)
                if hk is None or vk is None:
                    m.skipped[f"{hv} | {vv}"] = m.skipped.get(f"{hv} | {vv}", 0) + int(n)
                    continue
                m.counts[(hk, vk)] = m.counts.get((hk, vk), 0) + int(n)
                m.areas[(hk, vk)] = m.areas.get((hk, vk), 0.0) + float(a)
                m.orig.setdefault((hk, vk), set()).add((orig_str(hv), orig_str(vv)))
        else:
            f = cfg["field"]
            m.h_title = m.v_title = f
            first_v = cfg.get("first_axis", "v") == "v"
            m.v_title = f"{f}: {'1-я' if first_v else '2-я'} часть"
            m.h_title = f"{f}: {'2-я' if first_v else '1-я'} часть"
            agg = df.groupby(f, dropna=False)[AREA_FIELD].agg(["size", "sum"])
            for val, (n, area_m2) in agg.iterrows():
                pr = parse_code(val, cfg["method"], int(cfg.get("k", 1)), cfg.get("sep", "_"))
                if pr is None:
                    kk = str(val)
                    m.skipped[kk] = m.skipped.get(kk, 0) + int(n)
                    continue
                p1, p2 = pr
                vk, hk = (p1, p2) if first_v else (p2, p1)
                m.counts[(hk, vk)] = m.counts.get((hk, vk), 0) + int(n)
                m.areas[(hk, vk)] = m.areas.get((hk, vk), 0.0) + float(area_m2)
                m.codes.setdefault((hk, vk), set()).add(raw_code(val))
                m.orig.setdefault((hk, vk), set()).add((orig_str(val),))
        m.h_keys = sort_keys({h for h, _ in m.counts})
        m.v_keys = sort_keys({v for _, v in m.counts})
        if not m.h_keys or not m.v_keys:
            raise ValueError("Не найдено ни одного сочетания классов — проверьте поля и способ разбора.")
        m.hue_side = "h" if len(m.h_keys) >= len(m.v_keys) else "v"
        m.apply_palette(m.palette)
        m._base = copy.deepcopy((m.counts, m.areas, m.codes, m.orig))
        return m

    # ---------- перекодировка типов ----------
    def cell_from_text(self, txt):
        """Код, введённый пользователем → (ячейка (hk, vk), кортеж значений полей).
        Одно поле: код целиком (разбирается так же, как поле). Два поля: «строка столбец»
        через пробел, «;», «|», «/» или «,»."""
        txt = str(txt).strip()
        if not txt:
            raise ValueError("Код не задан.")
        src = self.source
        if src.get("mode") == "two":
            parts = [x for x in re.split(r"[\s;|/,]+", txt) if x]
            if len(parts) != 2:
                raise ValueError(f"«{txt}»: нужны два значения — строка и столбец "
                                 f"({src.get('v_field')} и {src.get('h_field')}), например «1 101».")
            vv, hv = parts
            hk, vk = norm_key(hv), norm_key(vv)
            if hk is None or vk is None:
                raise ValueError(f"«{txt}»: пустое значение.")
            return (hk, vk), (hv, vv)
        pr = parse_code(txt, src.get("method"), int(src.get("k", 1)), src.get("sep", "_"))
        if pr is None:
            raise ValueError(f"«{txt}» не разбирается выбранным способом "
                             f"({dict(PARSE_METHODS).get(src.get('method'), '')}).")
        p1, p2 = pr
        vk, hk = (p1, p2) if src.get("first_axis", "v") == "v" else (p2, p1)
        return (hk, vk), (txt,)

    def cell_text(self, hk, vk):
        """Код ячейки для подстановки в поля перекодировки."""
        if self.source.get("mode") == "two":
            return f"{vk} {hk}"
        if (hk, vk) in self.cell_code:
            return self.cell_code[(hk, vk)]
        codes = sorted(self.codes.get((hk, vk), []))
        if codes:
            return codes[0]
        a, b = (vk, hk) if self.source.get("first_axis", "v") == "v" else (hk, vk)
        return f"{a}{self.source.get('sep', '_')}{b}" if self.source.get("method") == "sep" else f"{a}{b}"

    def recode(self, old_txt, new_txt):
        """Перенести сведения ячейки «старый код» в ячейку «новый код» (слияние или
        перенумерация). → описание шага."""
        old, _ = self.cell_from_text(old_txt)
        new, new_vals = self.cell_from_text(new_txt)
        if old == new:
            raise ValueError("Старый и новый коды указывают на одну ячейку.")
        if old not in self.counts:
            raise ValueError(f"Код «{old_txt}» в матрице отсутствует (нет полигонов).")
        op = {"old": old, "new": new, "old_txt": str(old_txt).strip(),
              "new_txt": str(new_txt).strip(), "new_vals": new_vals}
        self.recodes.append(op)
        self.redo_stack.clear()
        self.undo_floor = max(self.undo_floor, len(self.recodes) - UNDO_LIMIT)
        self._recompute()
        return op

    def can_undo(self):
        return len(self.recodes) > self.undo_floor

    def undo(self):
        if not self.can_undo():
            return None
        op = self.recodes.pop()
        self.redo_stack.append(op)
        self._recompute()
        return op

    def redo(self):
        if not self.redo_stack:
            return None
        op = self.redo_stack.pop()
        self.recodes.append(op)
        self._recompute()
        return op

    def replay(self, pairs):
        """Повторить перекодировки [(старый, новый), ...]; неприменимые пропускаются."""
        bad = []
        for o, n in pairs:
            try:
                self.recode(o, n)
            except ValueError as ex:
                bad.append(f"{o} → {n}: {ex}")
        self.redo_stack.clear()
        return bad

    def _recompute(self):
        """Пересобрать матрицу: исходные сведения слоя + все шаги перекодировки по порядку."""
        counts, areas, codes, orig = copy.deepcopy(self._base)
        final, cell_code = {}, {}
        for op in self.recodes:
            o, n = op["old"], op["new"]
            if o not in counts:
                op["n"], op["area"] = 0, 0.0
                continue
            op["n"], op["area"] = counts[o], areas.get(o, 0.0)
            counts[n] = counts.get(n, 0) + counts.pop(o)
            areas[n] = areas.get(n, 0.0) + areas.pop(o, 0.0)
            codes.setdefault(n, set()).update(codes.pop(o, set()))
            moved = orig.pop(o, set())
            for t in moved:
                final[t] = tuple(op["new_vals"])
            orig.setdefault(n, set()).update(moved)
            if self.source.get("mode") != "two":
                cell_code[n] = op["new_txt"]
            cell_code.pop(o, None)
            self.manual.pop(o, None)
        self.counts, self.areas, self.codes, self.orig = counts, areas, codes, orig
        self.final, self.cell_code = final, cell_code
        self.h_keys = sort_keys({h for h, _ in counts})
        self.v_keys = sort_keys({v for _, v in counts})
        miss = [k for k in self.hue_keys if k not in self.base]
        if miss:                                    # новые категории цветовой оси — цвет из палитры
            used = set(self.base.values())
            cols = [c for c in palette_colors(self.palette, len(self.hue_keys) + len(self.base))
                    if c not in used]
            for k, c in zip(miss, cols + ["#888888"] * len(miss)):
                self.base[k] = c

    def set_names(self, h_names=None, v_names=None):
        if h_names is not None:
            self.h_names = dict(h_names)
        if v_names is not None:
            self.v_names = dict(v_names)

    # ---------- оси ----------
    @property
    def hue_keys(self):
        return self.h_keys if self.hue_side == "h" else self.v_keys

    @property
    def tone_keys(self):
        return self.v_keys if self.hue_side == "h" else self.h_keys

    def name(self, side, key):
        names = self.h_names if side == "h" else self.v_names
        return names.get(key, "")

    def label(self, side, key):
        """Подпись строки/столбца: название (по умолчанию; номер, если названия нет),
        «номер — название» или номер."""
        n = self.name(side, key)
        if self.label_mode == "num" or not n:
            return str(key)
        if self.label_mode == "num_name":
            return f"{key} — {n}"
        return n

    def area(self, hk, vk):
        return self.areas.get((hk, vk), 0.0) / AREA_UNITS[self.area_unit]

    def set_hue_side(self, side):
        if side != self.hue_side:
            self.hue_side = side
            self.override.clear()
            self.apply_palette(self.palette)

    def apply_palette(self, name):
        self.palette = name
        cols = palette_colors(name, len(self.hue_keys))
        self.base = {k: c for k, c in zip(self.hue_keys, cols)}

    # ---------- цвета ----------
    def tone_settings(self, hue_key):
        s = dict(self.tone)
        s.update(self.override.get(hue_key, {}))
        return s

    def auto_color(self, hk, vk):
        hue_key, tone_key = (hk, vk) if self.hue_side == "h" else (vk, hk)
        tk_ = self.tone_keys
        s = self.tone_settings(hue_key)
        return tone_color(self.base[hue_key], tk_.index(tone_key), len(tk_), s["mode"],
                          float(s["strength"]), bool(s["reverse"]), int(self.tone.get("seed", 1)),
                          self.hue_keys.index(hue_key))

    def color(self, hk, vk):
        return self.manual.get((hk, vk)) or self.auto_color(hk, vk)

    def present(self, hk, vk):
        return (hk, vk) in self.counts

    # ---------- состояние ----------
    def state(self):
        return {"hue_side": self.hue_side, "base": self.base, "tone": self.tone,
                "override": self.override, "palette": self.palette,
                "manual": [[h, v, c] for (h, v), c in self.manual.items()],
                "hide_absent": self.hide_absent, "label_mode": self.label_mode,
                "area_unit": self.area_unit,
                "recodes": [[op["old_txt"], op["new_txt"]] for op in self.recodes]}

    def load_state(self, st):
        """Применить сохранённую раскраску; категории, которых нет, пропускаются.
        Перекодировки из сохранённого файла применяются, если в модели их ещё нет."""
        self.load_bad = []
        if st.get("recodes") and not self.recodes:
            self.load_bad = self.replay(st["recodes"])
        if st.get("hue_side") in ("h", "v"):
            self.hue_side = st["hue_side"]
        self.apply_palette(st.get("palette", self.palette))
        for k, c in st.get("base", {}).items():
            if k in self.base and is_hex(c):
                self.base[k] = norm_hex(c)
        self.tone.update({k: v for k, v in st.get("tone", {}).items() if k in self.tone})
        self.override = {k: v for k, v in st.get("override", {}).items() if k in self.base}
        hs, vs = set(self.h_keys), set(self.v_keys)
        self.manual = {(h, v): norm_hex(c) for h, v, c in st.get("manual", [])
                       if h in hs and v in vs and is_hex(c)}
        self.hide_absent = bool(st.get("hide_absent", self.hide_absent))
        if st.get("label_mode") in dict(LABEL_MODES):
            self.label_mode = st["label_mode"]
        if st.get("area_unit") in AREA_UNITS:
            self.area_unit = st["area_unit"]


# ===========================================================================
# 3. ЦВЕТ
# ===========================================================================

def is_hex(s):
    return isinstance(s, str) and re.fullmatch(r"#?[0-9a-fA-F]{6}", s.strip()) is not None


def norm_hex(s):
    s = s.strip().lstrip("#")
    return "#" + s.lower()


def hex2rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) / 255.0 for i in (0, 2, 4))


def rgb2hex(rgb):
    return "#" + "".join(f"{max(0, min(255, int(round(c * 255)))):02x}" for c in rgb)


def mix(c1, c2, t):
    return tuple(a + (b - a) * t for a, b in zip(c1, c2))


def luminance(h):
    r, g, b = hex2rgb(h)
    return 0.299 * r + 0.587 * g + 0.114 * b


def text_color(h):
    return "#000000" if luminance(h) > 0.55 else "#ffffff"


def palette_colors(name, n):
    base = PALETTES.get(name)
    if base:
        if n <= len(base):
            return base[:n]
        # не хватает — продолжаем по золотому сечению
        return base + palette_colors("Золотое сечение (оттенок)", n - len(base))
    if name.startswith("Радуга"):
        return [rgb2hex(colorsys.hsv_to_rgb((i / max(n, 1)) * 0.83, 0.78, 0.88)) for i in range(n)]
    out = []
    for i in range(n):
        h = (0.07 + i * 0.6180339887) % 1.0
        v = (0.90, 0.72, 0.82)[i % 3]
        out.append(rgb2hex(colorsys.hsv_to_rgb(h, 0.75, v)))
    return out


_RANDOM_CACHE = {}


def _random_scale(base_hex, n, strength, seed, hue_index):
    key = (base_hex, n, round(strength, 4), seed, hue_index)
    if key in _RANDOM_CACHE:
        return _RANDOM_CACHE[key]
    rng = random.Random(seed * 100003 + hue_index * 7919 + n)
    h, l, s = colorsys.rgb_to_hls(*hex2rgb(base_hex))
    span = 0.32 * strength
    lo, hi = max(0.12, l - span), min(0.90, l + span)
    ls = [lo + (hi - lo) * i / max(n - 1, 1) for i in range(n)]
    rng.shuffle(ls)
    out = []
    for i in range(n):
        hh = (h + rng.uniform(-1, 1) * (12.0 / 360.0) * strength) % 1.0
        ss = min(1.0, max(0.15, s * rng.uniform(1 - 0.35 * strength, 1.0)))
        out.append(rgb2hex(colorsys.hls_to_rgb(hh, ls[i], ss)))
    if n >= 1:
        out[0] = base_hex if n == 1 else out[0]
    _RANDOM_CACHE[key] = out
    return out


def tone_color(base_hex, i, n, mode, strength=0.75, reverse=False, seed=1, hue_index=0):
    """Цвет i-й категории тоновой оси (из n) для базового цвета."""
    if n <= 1:
        return base_hex
    t = i / (n - 1)
    if reverse:
        t = 1.0 - t
    s = max(0.0, min(1.0, strength))
    c = hex2rgb(base_hex)
    if mode == "tint":
        return rgb2hex(mix(c, (1, 1, 1), s * t))
    if mode == "shade":
        return rgb2hex(mix(c, (0, 0, 0), s * t))
    if mode == "grey":
        g = 0.299 * c[0] + 0.587 * c[1] + 0.114 * c[2]
        return rgb2hex(mix(c, (g, g, g), s * t))
    if mode == "tint_shade":
        u = 2 * t - 1
        return rgb2hex(mix(c, (1, 1, 1), s * -u) if u < 0 else mix(c, (0, 0, 0), s * u * 0.85))
    if mode == "lightness":
        h, l, sat = colorsys.rgb_to_hls(*c)
        top, bot = 0.5 + 0.40 * s, 0.5 - 0.36 * s
        return rgb2hex(colorsys.hls_to_rgb(h, top + (bot - top) * t, sat))
    if mode == "random":
        idx = int(round(t * (n - 1)))
        return _random_scale(base_hex, n, s, seed, hue_index)[idx]
    raise ValueError(f"Неизвестный способ тона: {mode}")


# ===========================================================================
# 4. ЗАПИСЬ XLSX
# ===========================================================================

def save_xlsx(m, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb = Workbook()
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    def fill(h):
        return PatternFill("solid", fgColor=h.lstrip("#").upper())

    def headers(ws, rows_per):
        """Заголовки: строка 1 — столбцы, колонка A — строки (rows_per строк Excel на строку матрицы)."""
        note = f"\nВ ячейке: число полигонов / площадь, {m.area_unit}"
        ws.cell(1, 1, f"Строки: {m.v_title}\nСтолбцы: {m.h_title}{note}").font = bold
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="bottom")
        for j, hk in enumerate(m.h_keys):
            c = ws.cell(1, j + 2, m.label("h", hk))
            c.alignment = Alignment(text_rotation=90, wrap_text=True, horizontal="center", vertical="bottom")
            c.font = bold
            c.border = border
            if m.hue_side == "h":
                c.fill = fill(m.base[hk])
                c.font = Font(bold=True, color=text_color(m.base[hk]).lstrip("#").upper())
            ws.column_dimensions[get_column_letter(j + 2)].width = 11
        for i, vk in enumerate(m.v_keys):
            r0 = 2 + i * rows_per
            if rows_per > 1:
                ws.merge_cells(start_row=r0, start_column=1, end_row=r0 + rows_per - 1, end_column=1)
            c = ws.cell(r0, 1, m.label("v", vk))
            c.font = bold
            c.alignment = Alignment(wrap_text=True, vertical="center")
            for k in range(rows_per):
                ws.cell(r0 + k, 1).border = border
                if m.hue_side == "v":
                    ws.cell(r0 + k, 1).fill = fill(m.base[vk])
            if m.hue_side == "v":
                c.font = Font(bold=True, color=text_color(m.base[vk]).lstrip("#").upper())
        ws.column_dimensions["A"].width = 42
        ws.row_dimensions[1].height = 150
        ws.freeze_panes = "B2"

    # ---------- Матрица: число полигонов / площадь ----------
    ws = wb.active
    ws.title = "Матрица"
    headers(ws, 2)
    afmt = AREA_FMT[m.area_unit]
    for i, vk in enumerate(m.v_keys):
        r0 = 2 + i * 2
        ws.row_dimensions[r0].height = 15
        ws.row_dimensions[r0 + 1].height = 15
        for j, hk in enumerate(m.h_keys):
            top, bot = ws.cell(r0, j + 2), ws.cell(r0 + 1, j + 2)
            top.alignment = Alignment(horizontal="center", vertical="bottom")
            bot.alignment = Alignment(horizontal="center", vertical="top")
            if m.hide_absent and not m.present(hk, vk):
                top.border = Border(left=thin, right=thin, top=thin)
                bot.border = Border(left=thin, right=thin, bottom=thin)
                continue
            h = m.color(hk, vk)
            tc = text_color(h).lstrip("#").upper()
            line = Side(style="thin", color=tc)          # горизонтальная черта между числом и площадью
            top.border = Border(left=thin, right=thin, top=thin, bottom=line)
            bot.border = Border(left=thin, right=thin, bottom=thin, top=line)
            for c in (top, bot):
                c.fill = fill(h)
                c.font = Font(size=9, color=tc)
            if m.present(hk, vk):
                top.value = m.counts[(hk, vk)]
                top.number_format = "0"
                bot.value = round(m.area(hk, vk), 6)
                bot.number_format = afmt

    # ---------- Классы ----------
    wc = wb.create_sheet("Классы")
    one = m.source.get("mode") == "one"
    head = (["Код"] if one else []) + [
        f"Горизонталь ({m.h_title})", "Название (горизонталь)",
        f"Вертикаль ({m.v_title})", "Название (вертикаль)",
        "HEX", "R", "G", "B", "Объектов", f"Площадь, {m.area_unit}", "Ручная правка",
        "Объединено из (исходные значения)"]
    for j, t in enumerate(head):
        c = wc.cell(1, j + 1, t)
        c.font = bold
    r = 2
    for vk in m.v_keys:
        for hk in m.h_keys:
            if not m.present(hk, vk):
                continue
            h = m.color(hk, vk)
            rgb = [int(round(x * 255)) for x in hex2rgb(h)]
            src_vals = sorted(QGIS_DELIM.join(t) if len(t) > 1 else t[0]
                              for t in m.orig.get((hk, vk), set()))
            merged = any(t in m.final for t in m.orig.get((hk, vk), set()))
            code = m.cell_code.get((hk, vk)) or ", ".join(sorted(m.codes.get((hk, vk), [])))
            row = ([_num(code)] if one else []) + [
                _num(hk), m.name("h", hk), _num(vk), m.name("v", vk),
                h.upper(), *rgb, m.counts[(hk, vk)], round(m.area(hk, vk), 6),
                "да" if (hk, vk) in m.manual else "",
                ", ".join(src_vals) if merged else ""]
            for j, v in enumerate(row):
                wc.cell(r, j + 1, v)
            wc.cell(r, head.index(f"Площадь, {m.area_unit}") + 1).number_format = afmt
            hc = wc.cell(r, head.index("HEX") + 1)
            hc.fill = fill(h)
            hc.font = Font(color=text_color(h).lstrip("#").upper())
            r += 1
    for j, w in enumerate([10] * (1 if one else 0) + [14, 34, 14, 34, 10, 6, 6, 6, 10, 14, 13, 30]):
        wc.column_dimensions[get_column_letter(j + 1)].width = w
    wc.freeze_panes = "A2"

    # ---------- Перекодировка ----------
    if m.recodes:
        wr = wb.create_sheet("Перекодировка")
        lab = "строка столбец" if not one else "код"
        for j, t in enumerate(["Шаг", f"Старый ({lab})", f"Новый ({lab})", "Старая ячейка",
                               "Новая ячейка", "Перенесено объектов", f"Перенесено, {m.area_unit}"]):
            wr.cell(1, j + 1, t).font = bold
        for i, op in enumerate(m.recodes):
            (oh, ov), (nh, nv) = op["old"], op["new"]
            row = [i + 1, _num(op["old_txt"]), _num(op["new_txt"]),
                   f"{m.label('v', ov)} — {m.label('h', oh)}",
                   f"{m.label('v', nv)} — {m.label('h', nh)}", op.get("n", 0),
                   round(op.get("area", 0.0) / AREA_UNITS[m.area_unit], 6)]
            for j, v in enumerate(row):
                wr.cell(i + 2, j + 1, v)
            wr.cell(i + 2, 7).number_format = afmt
        for j, w in enumerate([6, 16, 16, 50, 50, 12, 16]):
            wr.column_dimensions[get_column_letter(j + 1)].width = w
        wr.freeze_panes = "A2"

    # ---------- Цвета ----------
    wk = wb.create_sheet("Цвета")
    side = "h" if m.hue_side == "h" else "v"
    for j, t in enumerate(["Категория", "Название", "Базовый HEX", "Градиент", "Сила",
                           "Обратный порядок", "Свой градиент"]):
        wk.cell(1, j + 1, t).font = bold
    for i, k in enumerate(m.hue_keys):
        s = m.tone_settings(k)
        row = [_num(k), m.name(side, k), m.base[k].upper(), TONE_LABEL[s["mode"]],
               float(s["strength"]), "да" if s["reverse"] else "", "да" if k in m.override else ""]
        for j, v in enumerate(row):
            wk.cell(i + 2, j + 1, v)
        wk.cell(i + 2, 3).fill = fill(m.base[k])
        wk.cell(i + 2, 3).font = Font(color=text_color(m.base[k]).lstrip("#").upper())
    for j, w in enumerate([12, 36, 13, 38, 7, 16, 14]):
        wk.column_dimensions[get_column_letter(j + 1)].width = w

    # ---------- Параметры ----------
    wp = wb.create_sheet("Параметры")
    src = m.source
    meth = dict(PARSE_METHODS).get(src.get("method"), "")
    pars = [
        ("Дата", datetime.now().strftime("%d.%m.%Y %H:%M")),
        ("Слой", src.get("path", "")), ("Имя слоя", src.get("layer", "") or ""),
        ("Векторы", "два поля" if src.get("mode") == "two" else "одно поле"),
        ("Поле по горизонтали", src.get("h_field", "") if src.get("mode") == "two" else ""),
        ("Поле по вертикали", src.get("v_field", "") if src.get("mode") == "two" else ""),
        ("Поле кода", src.get("field", "") if src.get("mode") == "one" else ""),
        ("Разбор", meth), ("k", src.get("k", "") if src.get("mode") == "one" else ""),
        ("Разделитель", src.get("sep", "") if src.get("method") == "sep" else ""),
        ("Первая часть кода", ("по вертикали" if src.get("first_axis") == "v" else "по горизонтали")
         if src.get("mode") == "one" else ""),
        ("Названия по горизонтали (CSV)", src.get("names_h", "") or ""),
        ("Названия по вертикали (CSV)", src.get("names_v", "") or ""),
        ("Цвет задаёт ось", "горизонталь" if m.hue_side == "h" else "вертикаль"),
        ("Автопалитра", m.palette),
        ("Градиент по умолчанию", TONE_LABEL[m.tone["mode"]]),
        ("Сила", m.tone["strength"]), ("Обратный порядок", "да" if m.tone["reverse"] else "нет"),
        ("Зерно случайной шкалы", m.tone["seed"]),
        ("Сочетаний в слое", len(m.counts)),
        ("Размер матрицы", f"{len(m.v_keys)} × {len(m.h_keys)}"),
        ("Нераспознанных значений", sum(m.skipped.values())),
        ("Ручных правок ячеек", len(m.manual)),
        ("Шагов перекодировки", len(m.recodes)),
        ("Подписи строк и столбцов", dict(LABEL_MODES)[m.label_mode]),
        ("Единицы площади", m.area_unit),
        ("Площадь", m.source.get("area_note", "") or "по геометрии в СК слоя"),
        ("state_json", json.dumps(m.state(), ensure_ascii=False)),
    ]
    for i, (k, v) in enumerate(pars):
        wp.cell(i + 1, 1, k).font = bold
        wp.cell(i + 1, 2, v)
    wp.column_dimensions["A"].width = 32
    wp.column_dimensions["B"].width = 90
    wb.save(path)
    return path


# ===========================================================================
# 4б. ФАЙЛЫ СТИЛЕЙ: QGIS (.qml), ArcGIS Pro (.lyrx), ArcMap 10.8 (.py → .lyr)
# ===========================================================================

OUTLINE_RGB = (110, 110, 110)
QGIS_DELIM = "|"
ARCMAP_DELIM = ", "


def style_entries(m):
    """→ (поля, список классов): поля в порядке «строка, столбец» (или одно поле кода);
    класс = {values: [кортежи значений полей], label, hex, n}."""
    if m.source.get("mode") == "two":
        fields = [m.source["v_field"], m.source["h_field"]]
    else:
        fields = [m.source["field"]]
    out = []
    for vk in m.v_keys:
        for hk in m.h_keys:
            if not m.present(hk, vk):
                continue
            vals = sorted(m.orig.get((hk, vk), set()))
            if len(fields) == 2:
                vals = [(vv, hv) for hv, vv in vals]
            out.append({"values": vals, "label": f"{m.label('v', vk)} — {m.label('h', hk)}",
                        "hex": m.color(hk, vk), "n": m.counts[(hk, vk)]})
    return fields, out


def _rgb255(h):
    return [int(round(c * 255)) for c in hex2rgb(h)]


def write_qml(m, path):
    """QGIS: категоризированный стиль (одно поле — по полю, два — по выражению concat)."""
    import xml.etree.ElementTree as ET
    fields, entries = style_entries(m)
    if len(fields) == 1:
        attr = fields[0]
        vtype = {"integer": "long", "double": "double"}.get(m.field_types.get(fields[0]), "string")
    else:
        attr = f'concat("{fields[0]}", \'{QGIS_DELIM}\', "{fields[1]}")'
        vtype = "string"

    def symbol(parent, name, h):
        sym = ET.SubElement(parent, "symbol", {"name": name, "type": "fill", "alpha": "1",
                                               "clip_to_extent": "1", "force_rhr": "0",
                                               "is_animated": "0", "frame_rate": "10"})
        ET.SubElement(sym, "data_defined_properties")
        lay = ET.SubElement(sym, "layer", {"class": "SimpleFill", "enabled": "1", "locked": "0",
                                           "pass": "0"})
        r_, g_, b_ = _rgb255(h)
        props = {"border_width_map_unit_scale": "3x:0,0,0,0,0,0",
                 "color": f"{r_},{g_},{b_},255", "joinstyle": "bevel",
                 "offset": "0,0", "offset_map_unit_scale": "3x:0,0,0,0,0,0", "offset_unit": "MM",
                 "outline_color": ",".join(map(str, OUTLINE_RGB)) + ",255",
                 "outline_style": "solid", "outline_width": "0.26", "outline_width_unit": "MM",
                 "style": "solid"}
        opt = ET.SubElement(lay, "Option", {"type": "Map"})
        for k, v in props.items():
            ET.SubElement(opt, "Option", {"name": k, "type": "QString", "value": v})
        for k, v in props.items():           # для QGIS 3.0–3.14
            ET.SubElement(lay, "prop", {"k": k, "v": v})
        ET.SubElement(lay, "data_defined_properties")
        return sym

    root = ET.Element("qgis", {"version": "3.28.0", "styleCategories": "Symbology"})
    rend = ET.SubElement(root, "renderer-v2", {"type": "categorizedSymbol", "attr": attr,
                                               "forceraster": "0", "symbollevels": "0",
                                               "enableorderby": "0", "referencescale": "-1"})
    cats = ET.SubElement(rend, "categories")
    syms = ET.SubElement(rend, "symbols")
    i = 0
    for e in entries:
        for vals in e["values"]:
            val = vals[0] if len(fields) == 1 else QGIS_DELIM.join(vals)
            ET.SubElement(cats, "category", {"symbol": str(i), "value": val, "label": e["label"],
                                             "render": "true", "type": vtype, "uuid": f"c{i}"})
            symbol(syms, str(i), e["hex"])
            i += 1
    src = ET.SubElement(rend, "source-symbol")
    symbol(src, "0", "#cccccc")
    ET.SubElement(rend, "rotation")
    ET.SubElement(rend, "sizescale")
    ET.SubElement(root, "blendMode").text = "0"
    ET.SubElement(root, "featureBlendMode").text = "0"
    ET.SubElement(root, "layerOpacity").text = "1"
    ET.indent(root)
    with open(path, "w", encoding="utf-8") as f:
        f.write("<!DOCTYPE qgis PUBLIC 'http://mrcc.com/qgis.dtd' 'SYSTEM'>\n")
        f.write(ET.tostring(root, encoding="unicode"))
    return path


def _cim_color(h, alpha=100):
    return {"type": "CIMRGBColor", "values": _rgb255(h) + [alpha]}


def _cim_polygon(h):
    return {"type": "CIMSymbolReference", "symbol": {
        "type": "CIMPolygonSymbol",
        "symbolLayers": [
            {"type": "CIMSolidStroke", "enable": True, "capStyle": "Round", "joinStyle": "Round",
             "lineStyle3D": "Strip", "miterLimit": 10, "width": 0.4,
             "color": {"type": "CIMRGBColor", "values": list(OUTLINE_RGB) + [100]}},
            {"type": "CIMSolidFill", "enable": True, "color": _cim_color(h)}]}}


def _cim_data_connection(path, layer):
    p = os.path.abspath(path)
    ext = os.path.splitext(p)[1].lower()
    if ext == ".shp":
        return {"type": "CIMStandardDataConnection",
                "workspaceConnectionString": "DATABASE=" + os.path.dirname(p),
                "workspaceFactory": "Shapefile",
                "dataset": os.path.splitext(os.path.basename(p))[0],
                "datasetType": "esriDTFeatureClass"}
    if ext in (".gpkg", ".sqlite"):
        return {"type": "CIMStandardDataConnection",
                "workspaceConnectionString": "AUTHENTICATION_MODE=OSA;DATABASE=" + p,
                "workspaceFactory": "Sql",
                "dataset": "main." + (layer or os.path.splitext(os.path.basename(p))[0]),
                "datasetType": "esriDTFeatureClass"}
    return None


def write_lyrx(m, path):
    """ArcGIS Pro: слой с символикой «Уникальные значения» по одному или двум полям.
    Применяется инструментом «Применить символы из слоя» (Apply Symbology From Layer)
    или кнопкой «Импорт» на панели «Символы»; источник данных можно переназначить."""
    fields, entries = style_entries(m)
    name = os.path.splitext(os.path.basename(path))[0]
    classes = []
    for e in entries:
        classes.append({"type": "CIMUniqueValueClass", "label": e["label"], "patch": "Default",
                        "symbol": _cim_polygon(e["hex"]), "visible": True, "editable": True,
                        "values": [{"type": "CIMUniqueValue", "fieldValues": list(v)}
                                   for v in e["values"]]})
    renderer = {"type": "CIMUniqueValueRenderer", "fields": fields,
                "groups": [{"type": "CIMUniqueValueGroup", "classes": classes,
                            "heading": ", ".join(fields)}],
                "defaultLabel": "<прочие значения>", "defaultSymbol": _cim_polygon("#cccccc"),
                "defaultSymbolPatch": "Default", "useDefaultSymbol": False,
                "polygonSymbolColorTarget": "Fill"}
    uri = f"CIMPATH=map/{re.sub(r'[^0-9A-Za-z_]', '_', name)}.json"
    table = {"type": "CIMFeatureTable", "displayField": fields[0], "editable": True,
             "studyAreaSpatialRel": "esriSpatialRelUndefined", "searchOrder": "esriSearchOrderSpatial"}
    dc = _cim_data_connection(m.source.get("path", ""), m.source.get("layer"))
    if dc:
        table["dataConnection"] = dc
    layer = {"type": "CIMFeatureLayer", "name": name, "uRI": uri,
             "sourceModifiedTime": {"type": "TimeInstant"}, "useSourceMetadata": True,
             "layerType": "Operational", "showLegends": True, "visibility": True,
             "displayCacheType": "Permanent", "maxDisplayCacheAge": 5, "showPopups": True,
             "serviceLayerID": -1, "refreshRate": -1, "refreshRateUnit": "esriTimeUnitsSeconds",
             "autoGenerateFeatureTemplates": True, "featureElevationExpression": "0",
             "featureTable": table, "htmlPopupEnabled": True, "selectable": True,
             "featureCacheType": "Session", "renderer": renderer, "scaleSymbols": True,
             "snappable": True}
    doc = {"type": "CIMLayerDocument", "version": "2.6.0", "build": 24783,
           "layers": [uri], "layerDefinitions": [layer]}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False, indent=2)
    return path


def arcmap_layer_definition(m):
    """Определение слоя (layerDefinition) в формате веб-карты ArcGIS для
    arcpy.mapping.Layer.updateLayerFromJSON (ArcMap 10.x)."""
    fields, entries = style_entries(m)

    def sfs(h):
        return {"type": "esriSFS", "style": "esriSFSSolid", "color": _rgb255(h) + [255],
                "outline": {"type": "esriSLS", "style": "esriSLSSolid",
                            "color": list(OUTLINE_RGB) + [255], "width": 0.4}}
    infos = []
    for e in entries:
        for v in e["values"]:
            infos.append({"value": ARCMAP_DELIM.join(v), "label": e["label"], "description": "",
                          "symbol": sfs(e["hex"])})
    rend = {"type": "uniqueValue", "field1": fields[0], "defaultSymbol": sfs("#cccccc"),
            "defaultLabel": "Прочие значения", "uniqueValueInfos": infos}
    if len(fields) == 2:
        rend["field2"] = fields[1]
        rend["fieldDelimiter"] = ARCMAP_DELIM
    return {"drawingInfo": {"renderer": rend}}


ARCMAP_TEMPLATE = r'''# -*- coding: utf-8 -*-
"""
Матричная легенда для ArcMap 10.8 (создано 09_MatrixLegend.py, %(date)s).

ArcMap не позволяет создать файл .lyr вне программы, поэтому раскраска
применяется этим скриптом внутри ArcMap и затем сохраняется в .lyr:

1. Откройте ArcMap и добавьте полигональный слой с полями: %(fields_txt)s.
2. Откройте окно Python (Геообработка → Python) и выполните:
       execfile(r"%(script)s")
   или щёлкните в окне Python правой кнопкой → Загрузить… → этот файл → Enter.
3. Слой раскрасится; рядом будет сохранён файл слоя
       %(lyr)s
   Его можно применять к другим слоям: Свойства слоя → Символы → Импорт.

Если в таблице содержания несколько подходящих слоёв, укажите имя нужного
в LAYER_NAME ниже.
"""
from __future__ import print_function
import json
import os
import arcpy

LAYER_NAME = u""            # имя слоя в таблице содержания; пусто — первый слой с нужными полями
FIELDS = %(fields)s
OUT_LYR = r"%(lyr)s"
LAYER_DEFINITION = json.loads(r"""%(json)s""")


def _find_layer(mxd):
    for lyr in arcpy.mapping.ListLayers(mxd):
        if lyr.isGroupLayer or not lyr.isFeatureLayer:
            continue
        if LAYER_NAME and lyr.name != LAYER_NAME:
            continue
        names = [f.name.lower() for f in arcpy.ListFields(lyr)]
        if all(f.lower() in names for f in FIELDS):
            return lyr
    return None


def main():
    mxd = arcpy.mapping.MapDocument("CURRENT")
    lyr = _find_layer(mxd)
    if lyr is None:
        raise RuntimeError(u"Не найден слой с полями %%s" %% u", ".join(FIELDS))
    lyr.updateLayerFromJSON(json.dumps(LAYER_DEFINITION))
    arcpy.RefreshActiveView()
    arcpy.RefreshTOC()
    out = OUT_LYR
    if not os.path.isdir(os.path.dirname(out)):
        try:
            out = os.path.join(os.path.dirname(os.path.abspath(__file__)), os.path.basename(out))
        except NameError:
            out = os.path.join(arcpy.env.scratchFolder, os.path.basename(out))
    try:
        lyr.saveACopy(out, "10.8")
    except Exception:
        lyr.saveACopy(out)
    print(u"Слой «%%s» раскрашен, файл слоя: %%s" %% (lyr.name, out))


main()
'''


def write_arcmap_script(m, path, lyr_path):
    fields, _ = style_entries(m)
    js = json.dumps(arcmap_layer_definition(m), ensure_ascii=True)
    txt = ARCMAP_TEMPLATE % {
        "date": datetime.now().strftime("%d.%m.%Y %H:%M"),
        "fields_txt": ", ".join(fields), "fields": repr([str(f) for f in fields]) if all(
            f.isascii() for f in fields) else "[" + ", ".join(f"u{f!r}" for f in fields) + "]",
        "script": os.path.abspath(path), "lyr": os.path.abspath(lyr_path), "json": js}
    with open(path, "w", encoding="utf-8") as f:
        f.write(txt)
    return path


def write_recoded_layer(m, xlsx_path):
    """Копия исходного слоя с перекодированными значениями поля (полей) →
    <база>_recoded.gpkg рядом с XLSX. Без перекодировок — None."""
    if not m.final:
        return None
    import geopandas as gpd
    src = m.source
    gdf = gpd.read_file(src["path"], layer=src.get("layer") or None)
    fields = [src["h_field"], src["v_field"]] if src.get("mode") == "two" else [src["field"]]
    keys = list(zip(*[gdf[f].map(orig_str) for f in fields]))
    hit = [k in m.final for k in keys]
    for i, f in enumerate(fields):
        kind = m.field_types.get(f, "string")
        vals = []
        for k, h in zip(keys, hit):
            if not h:
                continue
            v = m.final[k][i]
            try:
                v = int(float(v)) if kind == "integer" else float(v) if kind == "double" else str(v)
            except ValueError:
                raise ValueError(f"Значение «{v}» нельзя записать в числовое поле {f}.")
            vals.append(v)
        gdf.loc[hit, f] = vals
    base = os.path.splitext(os.path.abspath(xlsx_path))[0]
    out = base + "_recoded.gpkg"
    lname = src.get("layer") or os.path.splitext(os.path.basename(src["path"]))[0]
    if os.path.exists(out):
        os.remove(out)
    gdf.to_file(out, layer=lname, driver="GPKG")
    log(f"Перекодированный слой: {out} (изменено объектов: {sum(hit)})")
    return out


def save_styles(m, xlsx_path):
    """Файлы стилей рядом с XLSX: <база>_qgis.qml, <база>_arcgispro.lyrx,
    <база>_arcmap108.py (создаёт <база>_arcmap108.lyr в ArcMap)."""
    base = os.path.splitext(os.path.abspath(xlsx_path))[0]
    return [write_qml(m, base + "_qgis.qml"),
            write_lyrx(m, base + "_arcgispro.lyrx"),
            write_arcmap_script(m, base + "_arcmap108.py", base + "_arcmap108.lyr")]


def _num(k):
    try:
        f = float(k)
        return int(f) if f.is_integer() else f
    except ValueError:
        return k


def load_state_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    if "Параметры" not in wb.sheetnames:
        raise ValueError("В файле нет листа «Параметры».")
    for row in wb["Параметры"].iter_rows(values_only=True):
        if row and row[0] == "state_json" and row[1]:
            return json.loads(row[1])
    raise ValueError("В листе «Параметры» нет сохранённой раскраски.")


def default_out_path(path, layer=None, root_dir=None):
    root_dir = root_dir or os.getcwd()
    stamp = datetime.now().strftime("%d%m%y_%H%M")
    mm = re.search(r"(_v\d{3})$", os.path.basename(os.path.dirname(os.path.abspath(path))))
    folder = os.path.join(root_dir, OUT_BASE_NAME, stamp + (mm.group(1) if mm else ""))
    base = os.path.splitext(os.path.basename(path))[0]
    if layer and layer != base:
        base = f"{base}_{layer}"
    return os.path.join(folder, f"{base}_legend.xlsx")


def _initial_dir(root_dir):
    for name in ("08_TOPOLOGY", "07_GENERALIZATION"):
        b = os.path.join(root_dir, name)
        if os.path.isdir(b):
            subs = [os.path.join(b, d) for d in os.listdir(b) if os.path.isdir(os.path.join(b, d))]
            return max(subs, key=os.path.getmtime) if subs else b
    return root_dir


# ===========================================================================
# 5. ОКНО
# ===========================================================================

class ScrollFrame:
    """Вертикально прокручиваемая панель (для левой панели настроек)."""

    def __init__(self, master, width=430):
        import tkinter as tk
        from tkinter import ttk
        self.outer = ttk.Frame(master)
        self.canvas = tk.Canvas(self.outer, width=width, highlightthickness=0)
        sb = ttk.Scrollbar(self.outer, orient="vertical", command=self.canvas.yview)
        self.inner = ttk.Frame(self.canvas)
        wid = self.canvas.create_window(0, 0, window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfigure(wid, width=e.width))
        self.canvas.configure(yscrollcommand=sb.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")
        self.outer.bind("<Enter>", self._bind)
        self.outer.bind("<Leave>", self._unbind)

    def _wheel(self, e):
        if getattr(e, "num", None) == 4:
            d = -1
        elif getattr(e, "num", None) == 5:
            d = 1
        else:
            d = -1 if e.delta > 0 else 1
        self.canvas.yview_scroll(d * 2, "units")

    def _bind(self, _):
        for ev in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            self.canvas.bind_all(ev, self._wheel)

    def _unbind(self, _):
        for ev in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            self.canvas.unbind_all(ev)


class App:
    CELL_W, CELL_H = 80, 40      # размер ячейки при масштабе 1
    RH, CH = 230, 150            # ширина заголовков строк / высота заголовков столбцов (px)

    def __init__(self, root_dir):
        import tkinter as tk
        from tkinter import ttk
        self.tk, self.ttk = tk, ttk
        self.root_dir = root_dir
        self.last_dir = _initial_dir(root_dir)
        self.df = None
        self.m = None
        self.z, self.ox, self.oy = 1.0, 0.0, 0.0
        self._redraw_pending = False
        self._drag = None

        w = self.win = tk.Tk()
        w.title("09 — Матричная легенда")
        w.geometry("1360x860")
        w.minsize(900, 560)
        w.attributes("-topmost", True)
        w.after(500, lambda: w.attributes("-topmost", False))

        pw = ttk.PanedWindow(w, orient="horizontal")
        pw.pack(fill="both", expand=True)
        self.sf = ScrollFrame(pw, width=440)
        pw.add(self.sf.outer, weight=0)
        right = ttk.Frame(pw)
        pw.add(right, weight=1)
        self._build_left(self.sf.inner)
        self._build_canvas(right)
        self.status = tk.StringVar(value="Выберите слой полигонов.")
        ttk.Label(w, textvariable=self.status, anchor="w", relief="sunken").pack(fill="x", side="bottom")

    # ------------------------------------------------------------------
    # левая панель
    # ------------------------------------------------------------------
    def _build_left(self, p):
        tk, ttk = self.tk, self.ttk
        pad = {"padx": 6, "pady": 3}

        # 1. слой
        f1 = ttk.LabelFrame(p, text="1. Слой полигонов")
        f1.pack(fill="x", padx=6, pady=4)
        self.v_path = tk.StringVar()
        self.v_layer = tk.StringVar()
        e = ttk.Entry(f1, textvariable=self.v_path)
        e.grid(row=0, column=0, columnspan=2, sticky="we", **pad)
        ttk.Button(f1, text="Обзор…", command=self.browse_layer).grid(row=0, column=2, **pad)
        ttk.Label(f1, text="Слой").grid(row=1, column=0, sticky="w", **pad)
        self.cb_layer = ttk.Combobox(f1, textvariable=self.v_layer, state="readonly", width=28)
        self.cb_layer.grid(row=1, column=1, columnspan=2, sticky="we", **pad)
        self.cb_layer.bind("<<ComboboxSelected>>", lambda e: self.load_attrs())
        f1.columnconfigure(1, weight=1)

        # 2. векторы
        f2 = ttk.LabelFrame(p, text="2. Векторы матрицы")
        f2.pack(fill="x", padx=6, pady=4)
        self.v_mode = tk.StringVar(value="two")
        ttk.Radiobutton(f2, text="В двух полях", value="two", variable=self.v_mode,
                        command=self._mode_changed).grid(row=0, column=0, sticky="w", **pad)
        ttk.Radiobutton(f2, text="В одном поле (код)", value="one", variable=self.v_mode,
                        command=self._mode_changed).grid(row=0, column=1, sticky="w", **pad)
        self.f_two = ttk.Frame(f2)
        self.f_two.grid(row=1, column=0, columnspan=2, sticky="we")
        self.v_hf, self.v_vf = tk.StringVar(), tk.StringVar()
        ttk.Label(self.f_two, text="По горизонтали").grid(row=0, column=0, sticky="w", **pad)
        self.cb_hf = ttk.Combobox(self.f_two, textvariable=self.v_hf, state="readonly", width=24)
        self.cb_hf.grid(row=0, column=1, sticky="we", **pad)
        ttk.Label(self.f_two, text="По вертикали").grid(row=1, column=0, sticky="w", **pad)
        self.cb_vf = ttk.Combobox(self.f_two, textvariable=self.v_vf, state="readonly", width=24)
        self.cb_vf.grid(row=1, column=1, sticky="we", **pad)
        ttk.Button(self.f_two, text="Поменять местами", command=self._swap_fields).grid(
            row=2, column=1, sticky="w", **pad)

        self.f_one = ttk.Frame(f2)
        self.v_f, self.v_meth = tk.StringVar(), tk.StringVar(value=PARSE_METHODS[0][1])
        self.v_k, self.v_sep = tk.StringVar(value="1"), tk.StringVar(value="_")
        self.v_first = tk.StringVar(value="по вертикали")
        ttk.Label(self.f_one, text="Поле кода").grid(row=0, column=0, sticky="w", **pad)
        self.cb_f = ttk.Combobox(self.f_one, textvariable=self.v_f, state="readonly", width=24)
        self.cb_f.grid(row=0, column=1, columnspan=3, sticky="we", **pad)
        ttk.Label(self.f_one, text="Разбор").grid(row=1, column=0, sticky="w", **pad)
        cb_m = ttk.Combobox(self.f_one, textvariable=self.v_meth, state="readonly", width=36,
                            values=[t for _, t in PARSE_METHODS])
        cb_m.grid(row=1, column=1, columnspan=3, sticky="we", **pad)
        ttk.Label(self.f_one, text="k цифр").grid(row=2, column=0, sticky="w", **pad)
        ttk.Spinbox(self.f_one, from_=1, to=8, textvariable=self.v_k, width=4,
                    command=self._preview).grid(row=2, column=1, sticky="w", **pad)
        ttk.Label(self.f_one, text="Разделитель").grid(row=2, column=2, sticky="e", **pad)
        ttk.Entry(self.f_one, textvariable=self.v_sep, width=4).grid(row=2, column=3, sticky="w", **pad)
        ttk.Label(self.f_one, text="1-я часть кода").grid(row=3, column=0, sticky="w", **pad)
        ttk.Combobox(self.f_one, textvariable=self.v_first, state="readonly", width=16,
                     values=["по вертикали", "по горизонтали"]).grid(row=3, column=1, columnspan=3,
                                                                     sticky="w", **pad)
        self.v_prev = tk.StringVar()
        ttk.Label(self.f_one, textvariable=self.v_prev, foreground="#555555",
                  wraplength=390, justify="left").grid(row=4, column=0, columnspan=4, sticky="w", **pad)
        for v in (self.v_f, self.v_meth, self.v_k, self.v_sep, self.v_first):
            v.trace_add("write", lambda *a: self._preview())
        self.f_one.columnconfigure(1, weight=1)

        # 3. названия
        f3 = ttk.LabelFrame(p, text="3. Названия классов (CSV: номер; название)")
        f3.pack(fill="x", padx=6, pady=4)
        self.v_nh, self.v_nv = tk.StringVar(), tk.StringVar()
        for r, (lab, var) in enumerate((("По горизонтали", self.v_nh), ("По вертикали", self.v_nv))):
            ttk.Label(f3, text=lab).grid(row=r, column=0, sticky="w", **pad)
            ttk.Entry(f3, textvariable=var, width=22).grid(row=r, column=1, sticky="we", **pad)
            ttk.Button(f3, text="Обзор…", width=7,
                       command=lambda v=var: self.browse_csv(v)).grid(row=r, column=2, **pad)
            ttk.Button(f3, text="×", width=2, command=lambda v=var: v.set("")).grid(row=r, column=3, **pad)
        ttk.Label(f3, text="Подписи осей").grid(row=2, column=0, sticky="w", **pad)
        self.v_lab = tk.StringVar(value=LABEL_MODES[0][1])
        cbl = ttk.Combobox(f3, textvariable=self.v_lab, state="readonly", width=20,
                           values=[t for _, t in LABEL_MODES])
        cbl.grid(row=2, column=1, sticky="w", **pad)
        cbl.bind("<<ComboboxSelected>>", lambda e: self._labels_changed())
        ttk.Label(f3, text="Площадь в").grid(row=3, column=0, sticky="w", **pad)
        self.v_unit = tk.StringVar(value="га")
        cbu = ttk.Combobox(f3, textvariable=self.v_unit, state="readonly", width=8,
                           values=list(AREA_UNITS))
        cbu.grid(row=3, column=1, sticky="w", **pad)
        cbu.bind("<<ComboboxSelected>>", lambda e: self._labels_changed())
        f3.columnconfigure(1, weight=1)
        ttk.Button(p, text="Построить матрицу", command=self.build).pack(fill="x", padx=12, pady=6)

        # 4. перекодировка типов
        f7 = ttk.LabelFrame(p, text="4. Перекодировка типов (слияние, перенумерация)")
        f7.pack(fill="x", padx=6, pady=4)
        self.v_rc_hint = tk.StringVar(value="Постройте матрицу.")
        ttk.Label(f7, textvariable=self.v_rc_hint, foreground="#555555", justify="left",
                  wraplength=380).grid(row=0, column=0, columnspan=4, sticky="w", **pad)
        self.v_rc_old, self.v_rc_new = tk.StringVar(), tk.StringVar()
        ttk.Label(f7, text="Старый").grid(row=1, column=0, sticky="w", **pad)
        e_old = ttk.Entry(f7, textvariable=self.v_rc_old, width=12)
        e_old.grid(row=1, column=1, sticky="we", **pad)
        ttk.Label(f7, text="→ Новый").grid(row=1, column=2, sticky="w", **pad)
        e_new = ttk.Entry(f7, textvariable=self.v_rc_new, width=12)
        e_new.grid(row=1, column=3, sticky="we", **pad)
        e_old.bind("<Return>", lambda e: e_new.focus_set())
        e_new.bind("<Return>", lambda e: self.do_recode())
        fbr = ttk.Frame(f7)
        fbr.grid(row=2, column=0, columnspan=4, sticky="w")
        ttk.Button(fbr, text="Выполнить", command=self.do_recode).pack(side="left", padx=6, pady=2)
        self.b_undo = ttk.Button(fbr, text="Отменить", command=self.undo_recode, state="disabled")
        self.b_undo.pack(side="left", padx=2, pady=2)
        self.b_redo = ttk.Button(fbr, text="Повторить", command=self.redo_recode, state="disabled")
        self.b_redo.pack(side="left", padx=2, pady=2)
        self.lb_rc = tk.Listbox(f7, height=5, activestyle="none")
        self.lb_rc.grid(row=3, column=0, columnspan=4, sticky="we", **pad)
        self.v_rc_info = tk.StringVar(value="")
        ttk.Label(f7, textvariable=self.v_rc_info, foreground="#555555").grid(
            row=4, column=0, columnspan=4, sticky="w", **pad)
        f7.columnconfigure(1, weight=1)
        f7.columnconfigure(3, weight=1)

        # 4. цвет и тон
        f4 = ttk.LabelFrame(p, text="5. Цвет и тон")
        f4.pack(fill="x", padx=6, pady=4)
        self.v_hue = tk.StringVar(value="h")
        self.rb_h = ttk.Radiobutton(f4, text="Цвет — по горизонтали", value="h",
                                    variable=self.v_hue, command=self._hue_changed)
        self.rb_v = ttk.Radiobutton(f4, text="Цвет — по вертикали", value="v",
                                    variable=self.v_hue, command=self._hue_changed)
        self.rb_h.grid(row=0, column=0, columnspan=2, sticky="w", **pad)
        self.rb_v.grid(row=1, column=0, columnspan=2, sticky="w", **pad)
        ttk.Label(f4, text="Автопалитра").grid(row=2, column=0, sticky="w", **pad)
        self.v_pal = tk.StringVar(value=DEFAULT_PALETTE)
        ttk.Combobox(f4, textvariable=self.v_pal, state="readonly", width=30,
                     values=list(PALETTES)).grid(row=2, column=1, sticky="we", **pad)
        ttk.Button(f4, text="Применить автопалитру", command=self._apply_palette).grid(
            row=3, column=1, sticky="w", **pad)
        ttk.Separator(f4).grid(row=4, column=0, columnspan=2, sticky="we", pady=4)
        ttk.Label(f4, text="Тон по умолчанию").grid(row=5, column=0, sticky="w", **pad)
        self.v_tmode = tk.StringVar(value=TONE_LABEL["tint"])
        cbt = ttk.Combobox(f4, textvariable=self.v_tmode, state="readonly", width=34,
                           values=[t for _, t in TONE_MODES])
        cbt.grid(row=5, column=1, sticky="we", **pad)
        cbt.bind("<<ComboboxSelected>>", lambda e: self._tone_changed())
        ttk.Label(f4, text="Сила").grid(row=6, column=0, sticky="w", **pad)
        self.v_str = tk.DoubleVar(value=0.75)
        fs = ttk.Frame(f4)
        fs.grid(row=6, column=1, sticky="we", **pad)
        ttk.Scale(fs, from_=0.1, to=1.0, variable=self.v_str,
                  command=lambda e: self._tone_changed()).pack(side="left", fill="x", expand=True)
        self.v_str_lbl = tk.StringVar(value="0.75")
        ttk.Label(fs, textvariable=self.v_str_lbl, width=5).pack(side="left")
        self.v_rev = tk.BooleanVar(value=False)
        ttk.Checkbutton(f4, text="Обратный порядок тона", variable=self.v_rev,
                        command=self._tone_changed).grid(row=7, column=1, sticky="w", **pad)
        ttk.Label(f4, text="Зерно (случайная)").grid(row=8, column=0, sticky="w", **pad)
        self.v_seed = tk.StringVar(value="1")
        sp = ttk.Spinbox(f4, from_=1, to=9999, textvariable=self.v_seed, width=6,
                         command=self._tone_changed)
        sp.grid(row=8, column=1, sticky="w", **pad)
        sp.bind("<Return>", lambda e: self._tone_changed())
        f4.columnconfigure(1, weight=1)

        # 5. категории цветовой оси
        self.f5 = ttk.LabelFrame(p, text="6. Категории цветовой оси")
        self.f5.pack(fill="x", padx=6, pady=4)
        ttk.Label(self.f5, text="Щелчок по цветному квадрату — изменить цвет категории;\n"
                                "«Градиент» — свой способ тона для категории.",
                  foreground="#555555", justify="left").pack(anchor="w", padx=6)
        self.f_cats = ttk.Frame(self.f5)
        self.f_cats.pack(fill="x")

        # 6. ячейки и сохранение
        f6 = ttk.LabelFrame(p, text="7. Ячейки и сохранение")
        f6.pack(fill="x", padx=6, pady=4)
        self.v_hide = tk.BooleanVar(value=True)
        ttk.Checkbutton(f6, text="Не раскрашивать сочетания, которых нет в слое",
                        variable=self.v_hide, command=self._hide_changed).grid(
            row=0, column=0, columnspan=3, sticky="w", **pad)
        ttk.Button(f6, text="Сбросить ручные правки ячеек", command=self._reset_manual).grid(
            row=1, column=0, columnspan=3, sticky="w", **pad)
        ttk.Button(f6, text="Загрузить раскраску из XLSX…", command=self._load_state).grid(
            row=2, column=0, columnspan=3, sticky="w", **pad)
        self.v_styles = tk.BooleanVar(value=True)
        ttk.Checkbutton(f6, text="Стили: QGIS, ArcGIS Pro, ArcMap 10.8",
                        variable=self.v_styles).grid(row=4, column=0, columnspan=3, sticky="w", **pad)
        self.v_wlayer = tk.BooleanVar(value=False)
        ttk.Checkbutton(f6, text="Записать перекодированный слой (GPKG)",
                        variable=self.v_wlayer).grid(row=5, column=0, columnspan=3, sticky="w", **pad)
        ttk.Label(f6, text="Файл XLSX").grid(row=3, column=0, sticky="w", **pad)
        self.v_out = tk.StringVar()
        ttk.Entry(f6, textvariable=self.v_out, width=24).grid(row=3, column=1, sticky="we", **pad)
        ttk.Button(f6, text="Обзор…", width=7, command=self.browse_out).grid(row=3, column=2, **pad)
        f6.columnconfigure(1, weight=1)
        fb = ttk.Frame(p)
        fb.pack(fill="x", padx=6, pady=8)
        ttk.Button(fb, text="Сохранить XLSX", command=self.save).pack(side="left", padx=6)
        ttk.Button(fb, text="Закрыть", command=self.win.destroy).pack(side="right", padx=6)
        self._mode_changed()

    # ------------------------------------------------------------------
    # холст матрицы
    # ------------------------------------------------------------------
    def _build_canvas(self, p):
        tk, ttk = self.tk, self.ttk
        bar = ttk.Frame(p)
        bar.pack(fill="x")
        ttk.Button(bar, text="Вписать (Home)", command=self.fit).pack(side="left", padx=4, pady=3)
        ttk.Button(bar, text="Масштаб 1:1", command=self.zoom_1).pack(side="left", padx=4)
        ttk.Button(bar, text="+", width=3, command=lambda: self.zoom_center(1.25)).pack(side="left")
        ttk.Button(bar, text="−", width=3, command=lambda: self.zoom_center(0.8)).pack(side="left", padx=4)
        ttk.Label(bar, text="В ячейках:").pack(side="left", padx=(12, 2))
        self.v_celltxt = tk.StringVar(value=CELL_TEXT[0][1])
        cbc = ttk.Combobox(bar, textvariable=self.v_celltxt, state="readonly", width=15,
                           values=[t for _, t in CELL_TEXT])
        cbc.pack(side="left")
        cbc.bind("<<ComboboxSelected>>", lambda e: self.redraw())
        ttk.Label(bar, foreground="#555555",
                  text="ЛКМ — цвет, ПКМ — меню, перетаскивание — сдвиг, колесо — масштаб").pack(
            side="left", padx=10)
        self.cv = c = tk.Canvas(p, background="#f4f4f4", highlightthickness=0)
        c.pack(fill="both", expand=True)
        c.bind("<Configure>", lambda e: self.redraw())
        c.bind("<ButtonPress-1>", self._press)
        c.bind("<B1-Motion>", self._motion)
        c.bind("<ButtonRelease-1>", self._release)
        c.bind("<ButtonPress-2>", lambda e: self._press(e, pan_only=True))
        c.bind("<B2-Motion>", self._motion)
        c.bind("<ButtonRelease-2>", self._release)
        c.bind("<Button-3>", self._menu)
        c.bind("<MouseWheel>", self._wheel)
        c.bind("<Shift-MouseWheel>", lambda e: self._pan_wheel(e, horiz=True))
        c.bind("<Button-4>", self._wheel)
        c.bind("<Button-5>", self._wheel)
        c.bind("<Shift-Button-4>", lambda e: self._pan_wheel(e, horiz=True))
        c.bind("<Shift-Button-5>", lambda e: self._pan_wheel(e, horiz=True))
        c.bind("<Enter>", lambda e: c.focus_set())
        for key, dx, dy in (("<Left>", 60, 0), ("<Right>", -60, 0), ("<Up>", 0, 60), ("<Down>", 0, -60)):
            c.bind(key, lambda e, dx=dx, dy=dy: self.pan(dx, dy))
        c.bind("<Home>", lambda e: self.fit())

    # ---------- геометрия ----------
    def cw(self):
        return self.CELL_W * self.z

    def ch(self):
        return self.CELL_H * self.z

    def cell_at(self, x, y):
        """→ ('cell', j, i) | ('hhead', j) | ('vhead', i) | None."""
        if self.m is None:
            return None
        nj, ni = len(self.m.h_keys), len(self.m.v_keys)
        j = math.floor((x - self.RH - self.ox) / self.cw())
        i = math.floor((y - self.CH - self.oy) / self.ch())
        if x >= self.RH and y >= self.CH:
            if 0 <= j < nj and 0 <= i < ni:
                return ("cell", j, i)
        elif y < self.CH and x >= self.RH and 0 <= j < nj:
            return ("hhead", j)
        elif x < self.RH and y >= self.CH and 0 <= i < ni:
            return ("vhead", i)
        return None

    def redraw(self):
        if not self._redraw_pending:
            self._redraw_pending = True
            self.win.after_idle(self._draw)

    def _draw(self):
        self._redraw_pending = False
        c = self.cv
        c.delete("all")
        W, H = c.winfo_width(), c.winfo_height()
        if self.m is None:
            c.create_text(W / 2, H / 2, fill="#777777", font=("TkDefaultFont", 12),
                          text="Выберите слой и поля, затем нажмите «Построить матрицу».")
            return
        m = self.m
        cw, ch = self.cw(), self.ch()
        nj, ni = len(m.h_keys), len(m.v_keys)
        j0 = max(0, math.floor((-self.ox) / cw))
        j1 = min(nj, math.ceil((W - self.RH - self.ox) / cw))
        i0 = max(0, math.floor((-self.oy) / ch))
        i1 = min(ni, math.ceil((H - self.CH - self.oy) / ch))
        fsz = int(max(6, min(13, 8.5 * self.z)))
        mode_txt = next(k for k, t in CELL_TEXT if t == self.v_celltxt.get())
        show_txt = ch >= 14 and cw >= 40 and mode_txt != "none"
        two_lines = mode_txt == "count_area" and ch >= 28
        outline = "#9a9a9a" if cw > 8 else ""
        for i in range(i0, i1):
            vk = m.v_keys[i]
            y = self.CH + self.oy + i * ch
            for j in range(j0, j1):
                hk = m.h_keys[j]
                x = self.RH + self.ox + j * cw
                if m.hide_absent and not m.present(hk, vk):
                    c.create_rectangle(x, y, x + cw, y + ch, fill="#ffffff", outline=outline)
                    if cw > 10:
                        c.create_line(x + 3, y + 3, x + cw - 3, y + ch - 3, fill="#dddddd")
                    continue
                col = m.color(hk, vk)
                c.create_rectangle(x, y, x + cw, y + ch, fill=col, outline=outline)
                if (hk, vk) in m.manual and cw > 12:
                    s = min(10, cw / 4, ch / 2.5)
                    c.create_polygon(x + cw - s, y, x + cw, y, x + cw, y + s,
                                     fill=text_color(col), outline="")
                if show_txt:
                    tcol = text_color(col)
                    if m.present(hk, vk):
                        n_txt = str(m.counts[(hk, vk)])
                        a_txt = _fmt_area(m.area(hk, vk), m.area_unit)
                        if two_lines:
                            c.create_text(x + cw / 2, y + ch * 0.29, text=n_txt, fill=tcol,
                                          font=("TkDefaultFont", fsz))
                            c.create_line(x + cw * 0.18, y + ch / 2, x + cw * 0.82, y + ch / 2, fill=tcol)
                            c.create_text(x + cw / 2, y + ch * 0.71, text=a_txt, fill=tcol,
                                          font=("TkDefaultFont", fsz))
                        else:
                            c.create_text(x + cw / 2, y + ch / 2, text=n_txt, fill=tcol,
                                          font=("TkDefaultFont", fsz))
        # закреплённые заголовки
        hfs = int(max(7, min(11, 9 * self.z)))
        c.create_rectangle(0, 0, W, self.CH, fill="#e8e8e8", outline="")
        c.create_rectangle(0, 0, self.RH, H, fill="#e8e8e8", outline="")
        for j in range(j0, j1):
            hk = m.h_keys[j]
            x = self.RH + self.ox + j * cw
            if m.hue_side == "h":
                b = m.base[hk]
                c.create_rectangle(x, self.CH - 16, x + cw, self.CH, fill=b, outline="#777777")
            c.create_line(x, 0, x, self.CH, fill="#cccccc")
            if cw >= 7:
                f_ = int(max(6, min(hfs, cw * 0.75)))
                txt = _clip(m.label("h", hk), int((self.CH - 24) / (f_ * 0.62)))
                c.create_text(x + cw / 2, self.CH - 20, text=txt, angle=90, anchor="w",
                              font=("TkDefaultFont", f_), fill="#222222")
        for i in range(i0, i1):
            vk = m.v_keys[i]
            y = self.CH + self.oy + i * ch
            if m.hue_side == "v":
                c.create_rectangle(self.RH - 16, y, self.RH, y + ch, fill=m.base[vk], outline="#777777")
            c.create_line(0, y, self.RH, y, fill="#cccccc")
            if ch >= 7:
                f_ = int(max(6, min(hfs, ch * 0.75)))
                txt = _clip(m.label("v", vk), int((self.RH - 26) / (f_ * 0.62)))
                c.create_text(self.RH - 20, y + ch / 2, text=txt, anchor="e",
                              font=("TkDefaultFont", f_), fill="#222222")
        c.create_rectangle(0, 0, self.RH, self.CH, fill="#dcdcdc", outline="#bbbbbb")
        c.create_text(8, 8, anchor="nw", width=self.RH - 14, font=("TkDefaultFont", 9),
                      fill="#333333",
                      text=f"Строки: {m.v_title}\nСтолбцы: {m.h_title}\n\n"
                           f"{ni} × {nj}, сочетаний в слое: {len(m.counts)}\n"
                           f"Цвет: {'столбцы' if m.hue_side == 'h' else 'строки'}; "
                           f"тон: {'строки' if m.hue_side == 'h' else 'столбцы'}")

    # ---------- навигация ----------
    def _clamp(self):
        """Матрица не уходит за край окна целиком."""
        if self.m is None:
            return
        W = self.cv.winfo_width() - self.RH
        H = self.cv.winfo_height() - self.CH
        mw, mh = len(self.m.h_keys) * self.cw(), len(self.m.v_keys) * self.ch()
        keep = 60
        self.ox = min(max(self.ox, min(0, W - mw) - keep if mw > W else -mw + keep), max(W - keep, 0))
        self.oy = min(max(self.oy, min(0, H - mh) - keep if mh > H else -mh + keep), max(H - keep, 0))

    def pan(self, dx, dy):
        self.ox += dx
        self.oy += dy
        self._clamp()
        self.redraw()

    def _zoom_at(self, f, x, y):
        nz = max(0.08, min(8.0, self.z * f))
        f = nz / self.z
        px, py = x - self.RH, y - self.CH
        self.ox = px - (px - self.ox) * f
        self.oy = py - (py - self.oy) * f
        self.z = nz
        self._clamp()
        self.redraw()

    def zoom_center(self, f):
        self._zoom_at(f, (self.cv.winfo_width() + self.RH) / 2, (self.cv.winfo_height() + self.CH) / 2)

    def zoom_1(self):
        self.z, self.ox, self.oy = 1.0, 0.0, 0.0
        self.redraw()

    def view_readable(self):
        """Начальный вид: вписать, но не мельче читаемого масштаба (дальше — сдвиг)."""
        self.fit()
        if self.z < 0.75:
            self.z, self.ox, self.oy = 0.75, 0.0, 0.0
            self.redraw()

    def fit(self):
        if self.m is None:
            return
        W, H = self.cv.winfo_width() - self.RH - 10, self.cv.winfo_height() - self.CH - 10
        nj, ni = len(self.m.h_keys), len(self.m.v_keys)
        self.z = max(0.08, min(3.0, W / (nj * self.CELL_W), H / (ni * self.CELL_H)))
        self.ox = self.oy = 0.0
        self.redraw()

    def _wheel(self, e):
        up = (getattr(e, "num", None) == 4) or (getattr(e, "delta", 0) > 0)
        self._zoom_at(1.15 if up else 1 / 1.15, e.x, e.y)

    def _pan_wheel(self, e, horiz=False):
        up = (getattr(e, "num", None) == 4) or (getattr(e, "delta", 0) > 0)
        self.pan(80 if up else -80, 0)

    def _press(self, e, pan_only=False):
        self._drag = {"x": e.x, "y": e.y, "moved": False, "pan_only": pan_only}

    def _motion(self, e):
        d = self._drag
        if not d:
            return
        dx, dy = e.x - d["x"], e.y - d["y"]
        if not d["moved"] and abs(dx) + abs(dy) < 4:
            return
        d["moved"] = True
        d["x"], d["y"] = e.x, e.y
        self.pan(dx, dy)

    def _release(self, e):
        d, self._drag = self._drag, None
        if not d or d["moved"] or d["pan_only"]:
            return
        hit = self.cell_at(e.x, e.y)
        if not hit:
            return
        if hit[0] == "cell":
            self.edit_cell(self.m.h_keys[hit[1]], self.m.v_keys[hit[2]])
        elif hit[0] == "hhead" and self.m.hue_side == "h":
            self.edit_base(self.m.h_keys[hit[1]])
        elif hit[0] == "vhead" and self.m.hue_side == "v":
            self.edit_base(self.m.v_keys[hit[1]])

    def _menu(self, e):
        hit = self.cell_at(e.x, e.y)
        if not hit:
            return
        tk = self.tk
        m = self.m
        menu = tk.Menu(self.win, tearoff=0)
        if hit[0] == "cell":
            hk, vk = m.h_keys[hit[1]], m.v_keys[hit[2]]
            hue_key = hk if m.hue_side == "h" else vk
            n = m.counts.get((hk, vk), 0)
            menu.add_command(label=f"{m.label('v', vk)}  ×  {m.label('h', hk)}  ({n} объектов)",
                             state="disabled")
            menu.add_separator()
            menu.add_command(label="Изменить цвет ячейки…", command=lambda: self.edit_cell(hk, vk))
            menu.add_command(label="Ввести HEX ячейки…", command=lambda: self.edit_cell(hk, vk, hex_=True))
            menu.add_command(label="Сбросить ячейку",
                             command=lambda: (m.manual.pop((hk, vk), None), self.redraw()))
            menu.add_separator()
            code = m.cell_text(hk, vk)
            menu.add_command(label=f"Перекодировка: старый код = {code}",
                             command=lambda: self.v_rc_old.set(code), state="normal" if n else "disabled")
            menu.add_command(label=f"Перекодировка: новый код = {code}",
                             command=lambda: self.v_rc_new.set(code))
        else:
            hue_key = None
            if hit[0] == "hhead" and m.hue_side == "h":
                hue_key = m.h_keys[hit[1]]
            if hit[0] == "vhead" and m.hue_side == "v":
                hue_key = m.v_keys[hit[1]]
            if hue_key is None:
                return
        side = m.hue_side
        menu.add_separator()
        menu.add_command(label=f"Базовый цвет «{_clip(m.label(side, hue_key), 40)}»…",
                         command=lambda: self.edit_base(hue_key))
        menu.add_command(label="Ввести HEX базового цвета…",
                         command=lambda: self.edit_base(hue_key, hex_=True))
        menu.add_command(label="Градиент этой категории…", command=lambda: self.edit_gradient(hue_key))
        menu.add_separator()
        menu.add_command(label="Сбросить все ручные правки ячеек", command=self._reset_manual)
        menu.tk_popup(e.x_root, e.y_root)

    # ------------------------------------------------------------------
    # действия
    # ------------------------------------------------------------------
    def browse_layer(self):
        from tkinter import filedialog
        p = filedialog.askopenfilename(parent=self.win, title="Слой полигонов",
                                       initialdir=self.last_dir, filetypes=VECTOR_TYPES)
        if not p:
            return
        self.last_dir = os.path.dirname(p)
        self.v_path.set(p)
        lays = list_layers(p)
        self.cb_layer["values"] = lays
        self.v_layer.set(lays[0] if lays else "")
        self.load_attrs()

    def load_attrs(self):
        from tkinter import messagebox
        p = self.v_path.get().strip()
        if not p:
            return
        try:
            self.df = read_attrs(p, self.v_layer.get() or None)
        except Exception as ex:
            messagebox.showerror("Ошибка чтения", str(ex), parent=self.win)
            return
        cols = [c for c in self.df.columns if c != AREA_FIELD]
        for cb in (self.cb_hf, self.cb_vf, self.cb_f):
            cb["values"] = cols
        # угадываем поля
        low = {c.lower(): c for c in cols}
        guess = next((low[k] for k in ("code", "class_value", "class", "gridcode", "kod") if k in low),
                     cols[0] if cols else "")
        self.v_f.set(guess)
        if len(cols) >= 2:
            self.v_hf.set(cols[0])
            self.v_vf.set(cols[1])
        self.v_out.set(default_out_path(p, self.v_layer.get(), self.root_dir))
        self.status.set(f"Объектов: {len(self.df)}, полей: {len(cols)}.")
        self._preview()

    def browse_csv(self, var):
        from tkinter import filedialog
        p = filedialog.askopenfilename(parent=self.win, title="CSV с названиями классов",
                                       initialdir=self.last_dir, filetypes=CSV_TYPES)
        if p:
            self.last_dir = os.path.dirname(p)
            var.set(p)

    def browse_out(self):
        from tkinter import filedialog
        cur = self.v_out.get()
        p = filedialog.asksaveasfilename(parent=self.win, title="Файл легенды",
                                         defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
                                         initialdir=os.path.dirname(cur) if cur else self.last_dir,
                                         initialfile=os.path.basename(cur) if cur else "legend.xlsx")
        if p:
            self.v_out.set(p)

    def _mode_changed(self):
        if self.v_mode.get() == "two":
            self.f_one.grid_remove()
            self.f_two.grid(row=1, column=0, columnspan=2, sticky="we")
        else:
            self.f_two.grid_remove()
            self.f_one.grid(row=1, column=0, columnspan=2, sticky="we")
            self._preview()

    def _swap_fields(self):
        a, b = self.v_hf.get(), self.v_vf.get()
        self.v_hf.set(b)
        self.v_vf.set(a)

    def _parse_cfg(self):
        meth = next(k for k, t in PARSE_METHODS if t == self.v_meth.get())
        try:
            k = max(1, int(self.v_k.get()))
        except ValueError:
            k = 1
        return meth, k, self.v_sep.get()

    def _preview(self):
        if self.df is None or self.v_mode.get() != "one" or self.v_f.get() not in self.df.columns:
            self.v_prev.set("")
            return
        meth, k, sep = self._parse_cfg()
        vals = list(self.df[self.v_f.get()].dropna().unique()[:400])
        ok = [(v, parse_code(v, meth, k, sep)) for v in vals]
        good = [(v, pr) for v, pr in ok if pr]
        bad = [v for v, pr in ok if not pr]
        first_v = self.v_first.get() == "по вертикали"
        ex = "; ".join(f"{raw_code(v)} → {a} | {b}" for v, (a, b) in good[:4])
        t = f"Пример: {ex}" if ex else "Ни одно значение не разобрано."
        t += f"\n1-я часть → {'строки' if first_v else 'столбцы'}, 2-я → {'столбцы' if first_v else 'строки'}."
        if bad:
            t += f"\nНе разобрано (из первых {len(vals)}): {len(bad)}, напр. {', '.join(map(str, bad[:3]))}"
        self.v_prev.set(t)

    def build(self):
        from tkinter import messagebox
        if self.df is None:
            messagebox.showwarning("Нет данных", "Сначала выберите слой.", parent=self.win)
            return
        cfg = {"path": self.v_path.get(), "layer": self.v_layer.get(), "mode": self.v_mode.get(),
               "names_h": self.v_nh.get().strip(), "names_v": self.v_nv.get().strip()}
        if cfg["mode"] == "two":
            cfg.update(h_field=self.v_hf.get(), v_field=self.v_vf.get())
            if not cfg["h_field"] or not cfg["v_field"] or cfg["h_field"] == cfg["v_field"]:
                messagebox.showwarning("Поля", "Выберите два разных поля.", parent=self.win)
                return
        else:
            meth, k, sep = self._parse_cfg()
            cfg.update(field=self.v_f.get(), method=meth, k=k, sep=sep,
                       first_axis="v" if self.v_first.get() == "по вертикали" else "h")
        try:
            m = MatrixModel.from_table(self.df, cfg)
            m.set_names(read_names_csv(cfg["names_h"]) if cfg["names_h"] else {},
                        read_names_csv(cfg["names_v"]) if cfg["names_v"] else {})
        except Exception as ex:
            messagebox.showerror("Ошибка", str(ex), parent=self.win)
            return
        old = self.m
        m.palette = self.v_pal.get()
        m.apply_palette(m.palette)
        m.tone.update(self._tone_from_ui())
        m.hide_absent = self.v_hide.get()
        m.label_mode = next(k for k, t in LABEL_MODES if t == self.v_lab.get())
        m.area_unit = self.v_unit.get()
        m.source["area_note"] = self.df.attrs.get("area_note", "")
        same_src = old is not None and all(old.source.get(k) == m.source.get(k) for k in (
            "path", "layer", "mode", "field", "h_field", "v_field", "method", "k", "sep", "first_axis"))
        if same_src and old.recodes:
            bad = m.replay([[op["old_txt"], op["new_txt"]] for op in old.recodes])
            for b in bad:
                log(f"Перекодировка не применена: {b}")
        if old is not None and old.h_keys == m.h_keys and old.v_keys == m.v_keys:
            m.load_state(old.state())      # перестроение тех же осей — сохраняем раскраску
        self.m = m
        self.v_hue.set(m.hue_side)
        self._update_hue_labels()
        self._fill_cats()
        miss_h = [k for k in m.h_keys if k not in m.h_names] if m.h_names else []
        miss_v = [k for k in m.v_keys if k not in m.v_names] if m.v_names else []
        msg = (f"Матрица {len(m.v_keys)} × {len(m.h_keys)}, сочетаний в слое: {len(m.counts)}."
               + (f" Нераспознанных значений: {sum(m.skipped.values())}." if m.skipped else "")
               + (f" Нет названий: по горизонтали {len(miss_h)}" if miss_h else "")
               + (f", по вертикали {len(miss_v)}." if miss_v else ""))
        self.status.set(msg)
        log(msg)
        self._rc_refresh()
        self.win.after(50, self.view_readable)

    # ---------- перекодировка ----------
    def _rc_refresh(self):
        m = self.m
        if m is None:
            return
        if m.source.get("mode") == "two":
            self.v_rc_hint.set(f"Ячейка задаётся парой «строка столбец» ({m.source.get('v_field')} "
                               f"{m.source.get('h_field')}), например «1 400». Правый щелчок по "
                               "ячейке подставляет её код.")
        else:
            self.v_rc_hint.set(f"Ячейка задаётся кодом поля {m.source.get('field')}, например "
                               "1400 → 3400 (слияние) или 1400 → 1450 (перенумерация). "
                               "Правый щелчок по ячейке подставляет её код.")
        self.lb_rc.delete(0, "end")
        for i, op in enumerate(m.recodes):
            self.lb_rc.insert("end", f"{i + 1}. {op['old_txt']} → {op['new_txt']}   "
                                     f"({op.get('n', 0)} об., {_fmt_area(op.get('area', 0.0) / AREA_UNITS[m.area_unit], m.area_unit)})")
        self.lb_rc.see("end")
        self.b_undo.configure(state="normal" if m.can_undo() else "disabled")
        self.b_redo.configure(state="normal" if m.redo_stack else "disabled")
        self.v_rc_info.set(f"Шагов: {len(m.recodes)}; можно отменить: {len(m.recodes) - m.undo_floor} "
                           f"(не более {UNDO_LIMIT}); повторить: {len(m.redo_stack)}.")

    def _after_recode(self, msg):
        self.v_hue.set(self.m.hue_side)
        self._update_hue_labels()
        self._fill_cats()
        self._clamp()
        self.redraw()
        self._rc_refresh()
        self.status.set(msg)
        log(msg)

    def do_recode(self):
        from tkinter import messagebox
        if not self.m:
            messagebox.showwarning("Нет матрицы", "Сначала постройте матрицу.", parent=self.win)
            return
        try:
            op = self.m.recode(self.v_rc_old.get(), self.v_rc_new.get())
        except ValueError as ex:
            messagebox.showwarning("Перекодировка", str(ex), parent=self.win)
            return
        m = self.m
        self.v_rc_old.set("")
        self._after_recode(f"{op['old_txt']} → {op['new_txt']}: перенесено {op['n']} объектов, "
                           f"{_fmt_area(op['area'] / AREA_UNITS[m.area_unit], m.area_unit)}. "
                           f"Матрица {len(m.v_keys)} × {len(m.h_keys)}, сочетаний: {len(m.counts)}.")

    def undo_recode(self, *_):
        if self.m and self.m.can_undo():
            op = self.m.undo()
            self._after_recode(f"Отменено: {op['old_txt']} → {op['new_txt']}.")

    def redo_recode(self, *_):
        if self.m and self.m.redo_stack:
            op = self.m.redo()
            self._after_recode(f"Повторено: {op['old_txt']} → {op['new_txt']}.")

    def _update_hue_labels(self):
        if self.m:
            self.rb_h.configure(text=f"Цвет — по горизонтали ({len(self.m.h_keys)} кат.)")
            self.rb_v.configure(text=f"Цвет — по вертикали ({len(self.m.v_keys)} кат.)")

    def _tone_from_ui(self):
        mode = next(k for k, t in TONE_MODES if t == self.v_tmode.get())
        try:
            seed = int(self.v_seed.get())
        except ValueError:
            seed = 1
        s = round(float(self.v_str.get()), 2)
        self.v_str_lbl.set(f"{s:.2f}")
        return {"mode": mode, "strength": s, "reverse": bool(self.v_rev.get()), "seed": seed}

    def _tone_changed(self):
        t = self._tone_from_ui()
        if self.m:
            self.m.tone.update(t)
            self._fill_cats()
            self.redraw()

    def _hue_changed(self):
        if self.m:
            self.m.set_hue_side(self.v_hue.get())
            self._fill_cats()
            self.redraw()

    def _apply_palette(self):
        if self.m:
            self.m.apply_palette(self.v_pal.get())
            self._fill_cats()
            self.redraw()

    def _labels_changed(self):
        if self.m:
            self.m.label_mode = next(k for k, t in LABEL_MODES if t == self.v_lab.get())
            self.m.area_unit = self.v_unit.get()
            self._fill_cats()
            self.redraw()

    def _hide_changed(self):
        if self.m:
            self.m.hide_absent = self.v_hide.get()
            self.redraw()

    def _reset_manual(self):
        if self.m and self.m.manual:
            self.m.manual.clear()
            self.redraw()

    def _fill_cats(self):
        tk, ttk = self.tk, self.ttk
        for w in self.f_cats.winfo_children():
            w.destroy()
        if not self.m:
            return
        m = self.m
        side = m.hue_side
        self.f5.configure(text=f"6. Категории цветовой оси ({'горизонталь' if side == 'h' else 'вертикаль'}, "
                               f"{len(m.hue_keys)})")
        for r, k in enumerate(m.hue_keys):
            sw = tk.Label(self.f_cats, width=3, background=m.base[k], relief="raised", cursor="hand2")
            sw.grid(row=r, column=0, padx=(6, 4), pady=1, sticky="w")
            sw.bind("<Button-1>", lambda e, k=k: self.edit_base(k))
            ttk.Label(self.f_cats, text=_clip(m.label(side, k), 34)).grid(row=r, column=1, sticky="w")
            gl = (TONE_SHORT[m.override[k]["mode"]] if k in m.override else "по умолч.")
            ttk.Button(self.f_cats, text=f"Градиент: {gl}", width=20,
                       command=lambda k=k: self.edit_gradient(k)).grid(row=r, column=2, padx=4, pady=1,
                                                                        sticky="e")
        self.f_cats.columnconfigure(1, weight=1)

    def _ask_color(self, init, title):
        from tkinter import colorchooser
        res = colorchooser.askcolor(color=init, parent=self.win, title=title)
        return norm_hex(res[1]) if res and res[1] else None

    def _ask_hex(self, init, title):
        from tkinter import simpledialog, messagebox
        s = simpledialog.askstring(title, "HEX-код цвета (например, #3a7bd5):",
                                   initialvalue=init.upper(), parent=self.win)
        if s is None:
            return None
        if not is_hex(s):
            messagebox.showwarning("HEX", f"Неверный код цвета: {s}", parent=self.win)
            return None
        return norm_hex(s)

    def edit_cell(self, hk, vk, hex_=False):
        m = self.m
        cur = m.color(hk, vk)
        title = f"Ячейка {m.label('v', vk)} × {m.label('h', hk)}"
        new = self._ask_hex(cur, title) if hex_ else self._ask_color(cur, title)
        if new:
            if new == m.auto_color(hk, vk):
                m.manual.pop((hk, vk), None)
            else:
                m.manual[(hk, vk)] = new
            self.redraw()

    def edit_base(self, k, hex_=False):
        m = self.m
        title = f"Базовый цвет: {m.label(m.hue_side, k)}"
        new = self._ask_hex(m.base[k], title) if hex_ else self._ask_color(m.base[k], title)
        if new:
            m.base[k] = new
            self._fill_cats()
            self.redraw()

    def edit_gradient(self, k):
        tk, ttk = self.tk, self.ttk
        m = self.m
        s = m.tone_settings(k)
        top = tk.Toplevel(self.win)
        top.title(f"Градиент: {_clip(m.label(m.hue_side, k), 50)}")
        top.transient(self.win)
        top.resizable(False, False)
        pad = {"padx": 8, "pady": 4}
        v_own = tk.BooleanVar(value=k in m.override)
        v_mode = tk.StringVar(value=TONE_LABEL[s["mode"]])
        v_str = tk.DoubleVar(value=float(s["strength"]))
        v_rev = tk.BooleanVar(value=bool(s["reverse"]))
        ttk.Checkbutton(top, text="Свой градиент (иначе — по умолчанию)", variable=v_own).grid(
            row=0, column=0, columnspan=2, sticky="w", **pad)
        ttk.Label(top, text="Способ").grid(row=1, column=0, sticky="w", **pad)
        ttk.Combobox(top, textvariable=v_mode, state="readonly", width=36,
                     values=[t for _, t in TONE_MODES]).grid(row=1, column=1, **pad)
        ttk.Label(top, text="Сила").grid(row=2, column=0, sticky="w", **pad)
        ttk.Scale(top, from_=0.1, to=1.0, variable=v_str).grid(row=2, column=1, sticky="we", **pad)
        ttk.Checkbutton(top, text="Обратный порядок", variable=v_rev).grid(row=3, column=1, sticky="w", **pad)
        prev = tk.Canvas(top, height=26, highlightthickness=0, background="#ffffff")
        prev.grid(row=4, column=0, columnspan=2, sticky="we", **pad)

        def cur():
            return {"mode": next(a for a, t in TONE_MODES if t == v_mode.get()),
                    "strength": round(float(v_str.get()), 2), "reverse": bool(v_rev.get())}

        def show(*_):
            prev.delete("all")
            sset = cur() if v_own.get() else {kk: m.tone[kk] for kk in ("mode", "strength", "reverse")}
            n = len(m.tone_keys)
            w = max(prev.winfo_width(), 340) / max(n, 1)
            for i in range(n):
                col = tone_color(m.base[k], i, n, sset["mode"], sset["strength"], sset["reverse"],
                                 int(m.tone.get("seed", 1)), m.hue_keys.index(k))
                prev.create_rectangle(i * w, 0, (i + 1) * w, 26, fill=col, outline="")

        for v in (v_own, v_mode, v_str, v_rev):
            v.trace_add("write", show)

        def ok():
            if v_own.get():
                m.override[k] = cur()
            else:
                m.override.pop(k, None)
            top.destroy()
            self._fill_cats()
            self.redraw()

        fb = ttk.Frame(top)
        fb.grid(row=5, column=0, columnspan=2, sticky="e", **pad)
        ttk.Button(fb, text="Применить", command=ok).pack(side="left", padx=4)
        ttk.Button(fb, text="Отмена", command=top.destroy).pack(side="left")
        top.after(60, show)
        top.grab_set()

    def _load_state(self):
        from tkinter import filedialog, messagebox
        if not self.m:
            messagebox.showwarning("Нет матрицы", "Сначала постройте матрицу.", parent=self.win)
            return
        p = filedialog.askopenfilename(parent=self.win, title="Сохранённая легенда XLSX",
                                       initialdir=self.last_dir, filetypes=[("Excel", "*.xlsx")])
        if not p:
            return
        try:
            self.m.load_state(load_state_xlsx(p))
        except Exception as ex:
            messagebox.showerror("Ошибка", str(ex), parent=self.win)
            return
        t = self.m.tone
        self.v_hue.set(self.m.hue_side)
        self.v_tmode.set(TONE_LABEL[t["mode"]])
        self.v_str.set(t["strength"])
        self.v_rev.set(bool(t["reverse"]))
        self.v_seed.set(str(t["seed"]))
        self.v_pal.set(self.m.palette)
        self.v_hide.set(self.m.hide_absent)
        self.v_lab.set(dict(LABEL_MODES)[self.m.label_mode])
        self.v_unit.set(self.m.area_unit)
        self._update_hue_labels()
        self._fill_cats()
        self._clamp()
        self.redraw()
        self._rc_refresh()
        bad = getattr(self.m, "load_bad", [])
        self.status.set(f"Раскраска загружена из {os.path.basename(p)}."
                        + (f" Перекодировок не применено: {len(bad)}." if bad else ""))

    def save(self):
        from tkinter import messagebox
        if not self.m:
            messagebox.showwarning("Нет матрицы", "Сначала постройте матрицу.", parent=self.win)
            return
        out = self.v_out.get().strip() or default_out_path(self.v_path.get(), self.v_layer.get(),
                                                           self.root_dir)
        if not out.lower().endswith(".xlsx"):
            out += ".xlsx"
        try:
            save_xlsx(self.m, out)
            styles = save_styles(self.m, out) if self.v_styles.get() else []
            if self.v_wlayer.get():
                lp = write_recoded_layer(self.m, out)
                if lp:
                    styles.append(lp)
        except PermissionError:
            messagebox.showerror("Ошибка записи", f"Файл занят (открыт в Excel?):\n{out}", parent=self.win)
            return
        except Exception as ex:
            messagebox.showerror("Ошибка записи", str(ex), parent=self.win)
            return
        self.v_out.set(out)
        self.status.set(f"Сохранено: {out}" + (f" + файлы ({len(styles)})" if styles else ""))
        log(f"Сохранено: {out}")
        for sp in styles:
            log(f"Файл: {sp}")

    def run(self):
        self.win.bind("<Control-z>", self.undo_recode)
        self.win.bind("<Control-y>", self.redo_recode)
        self.redraw()
        self.win.mainloop()


def _fmt_area(v, unit):
    if unit == "м²":
        t = f"{v:,.0f}"
    elif unit == "км²":
        t = f"{v:,.3f}"
    else:
        t = f"{v:,.2f}"
    return t.replace(",", " ").replace(".", ",") + " " + unit


def _clip(s, n):
    s = str(s)
    return s if len(s) <= max(n, 4) else s[:max(n, 4) - 1] + "…"


# ===========================================================================
# 6. ЗАПУСК
# ===========================================================================

def main(argv=None):
    if argv is None:
        argv = sys.argv[1:]
    # Jupyter передаёт свои аргументы (-f kernel.json) — считаем их пустыми
    if argv and any(a.startswith("-f") or a.endswith(".json") for a in argv) and \
            not any(a.lower().endswith((".gpkg", ".shp", ".geojson", ".fgb", ".sqlite")) for a in argv):
        argv = []
    if not argv:
        log("09 - Матричная легенда. Рабочая папка: " + os.getcwd())
        try:
            app = App(os.getcwd())
        except Exception as e:
            log(f"Графический интерфейс недоступен ({e}). Используйте CLI: python 09_MatrixLegend.py --help")
            return
        app.run()
        return

    p = argparse.ArgumentParser(description="Матричная легенда (раскраска по умолчанию, без окна)")
    p.add_argument("path", help="слой полигонов")
    p.add_argument("--layer", default=None)
    g = p.add_mutually_exclusive_group(required=True)
    g.add_argument("--fields", nargs=2, metavar=("ГОРИЗ", "ВЕРТ"), help="два поля: по горизонтали и по вертикали")
    g.add_argument("--field", help="одно поле с кодом")
    p.add_argument("--parse", default="first:1", help="first:k | last:k | sep:СИМВОЛ (для --field)")
    p.add_argument("--first-axis", choices=["v", "h"], default="v", help="ось 1-й части кода")
    p.add_argument("--names-h", default=None, help="CSV названий по горизонтали")
    p.add_argument("--names-v", default=None, help="CSV названий по вертикали")
    p.add_argument("--hue", choices=["h", "v"], default=None, help="ось цвета (по умолчанию — бо́льшая)")
    p.add_argument("--palette", default=DEFAULT_PALETTE, choices=list(PALETTES))
    p.add_argument("--tone", default="tint", choices=[k for k, _ in TONE_MODES])
    p.add_argument("--strength", type=float, default=0.75)
    p.add_argument("--reverse", action="store_true")
    p.add_argument("--seed", type=int, default=1)
    p.add_argument("--show-absent", action="store_true", help="раскрашивать и отсутствующие сочетания")
    p.add_argument("--labels", choices=[k for k, _ in LABEL_MODES], default="name",
                   help="подписи строк/столбцов: name — название, num_name — номер — название, num — номер")
    p.add_argument("--area-unit", choices=list(AREA_UNITS), default="га")
    p.add_argument("--no-styles", action="store_true", help="не создавать файлы стилей QGIS / ArcGIS")
    p.add_argument("--recode", nargs="+", default=[], metavar="СТАРЫЙ=НОВЫЙ",
                   help="перекодировка типов по порядку: 1400=3400 2400=3400 (два поля: \"1 400=3 400\")")
    p.add_argument("--write-layer", action="store_true",
                   help="записать копию слоя с перекодированными значениями (<база>_recoded.gpkg)")
    p.add_argument("--out", default=None)
    a = p.parse_args(argv)

    cfg = {"path": a.path, "layer": a.layer, "names_h": a.names_h, "names_v": a.names_v}
    if a.fields:
        cfg.update(mode="two", h_field=a.fields[0], v_field=a.fields[1])
    else:
        meth, _, arg = a.parse.partition(":")
        if meth not in ("first", "last", "sep"):
            p.error("--parse: first:k, last:k или sep:СИМВОЛ")
        cfg.update(mode="one", field=a.field, method=meth, first_axis=a.first_axis,
                   k=int(arg) if meth != "sep" and arg else 1, sep=arg if meth == "sep" else "_")
    df = read_attrs(a.path, a.layer)
    m = MatrixModel.from_table(df, cfg)
    m.set_names(read_names_csv(a.names_h) if a.names_h else {},
                read_names_csv(a.names_v) if a.names_v else {})
    if a.hue:
        m.set_hue_side(a.hue)
    m.apply_palette(a.palette)
    m.tone.update(mode=a.tone, strength=a.strength, reverse=a.reverse, seed=a.seed)
    m.hide_absent = not a.show_absent
    m.label_mode, m.area_unit = a.labels, a.area_unit
    m.source["area_note"] = df.attrs.get("area_note", "")
    if m.source["area_note"]:
        log(m.source["area_note"])
    for rc in a.recode:
        o, sep_, n = rc.partition("=")
        if not sep_:
            p.error(f"--recode: ожидается СТАРЫЙ=НОВЫЙ, получено «{rc}»")
        op = m.recode(o, n)
        log(f"{op['old_txt']} → {op['new_txt']}: перенесено {op['n']} объектов")
    out = a.out or default_out_path(a.path, a.layer)
    save_xlsx(m, out)
    styles = [] if a.no_styles else save_styles(m, out)
    if a.write_layer:
        lp = write_recoded_layer(m, out)
        if lp:
            styles.append(lp)
    log(f"Матрица {len(m.v_keys)} × {len(m.h_keys)}, сочетаний: {len(m.counts)}, "
        f"нераспознанных значений: {sum(m.skipped.values())}")
    log(f"Сохранено: {out}")
    for sp in styles:
        log(f"Файл: {sp}")


if __name__ == "__main__":
    main()
